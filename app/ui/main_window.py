import csv
import json
import os
import shutil
import sys
import tempfile
import statistics
import subprocess
from collections import Counter
from datetime import datetime
from html import escape
from pathlib import Path
from typing import Dict, List, Optional

import cv2
import numpy as np
from PyQt5.QtCore import Qt, QSettings, QTimer
from PyQt5.QtGui import QColor, QFont, QPainter, QPen, QPixmap
from PyQt5.QtWidgets import (
    QApplication,
    QCheckBox,
    QDialog,
    QDoubleSpinBox,
    QFileDialog,
    QGridLayout,
    QGroupBox,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QProgressBar,
    QSizePolicy,
    QSpinBox,
    QTabWidget,
    QTableWidget,
    QTableWidgetItem,
    QTextEdit,
    QVBoxLayout,
    QWidget,
    QInputDialog,
    QFrame,
)

from app.backend import LocalBackend
from app.config import AppConfig
from app.services.detector_service import DetectorService
from app.ui.styles import button_style, card_style, main_window_styles, path_label_style, theme_palette
from app.utils.file_utils import clear_dir, ensure_dir, list_images, safe_copy, short_path
from app.utils.image_utils import cv_to_pixmap
from app.workers.benchmark_worker import BenchmarkWorker
from app.workers.detection_worker import DetectionMode, DetectionWorker

try:
    import psutil
except Exception:  # pragma: no cover - psutil is listed in requirements
    psutil = None


class MainWindow(QMainWindow):
    def __init__(self) -> None:
        super().__init__()
        self.config = AppConfig()
        self.settings = QSettings(self.config.organization, self.config.app_id)
        self.detector = DetectorService()
        self.project_root = Path(__file__).resolve().parents[2]
        backend_db = self.project_root / self.config.backend_dir_name / self.config.backend_db_name
        self.backend = LocalBackend(str(backend_db))
        self.backend_session_id: Optional[int] = None

        self.worker: Optional[DetectionWorker] = None
        self.benchmark_worker: Optional[BenchmarkWorker] = None
        self.current_mode = DetectionMode.NONE
        self.current_image_path: Optional[str] = None
        self.current_folder_images: List[str] = []
        self.current_video_path: Optional[str] = None
        self.current_camera_id: Optional[int] = None

        self.last_original_image = None
        self.last_clean_annotated_image = None
        self.last_annotated_image = None
        self.last_rows: List[Dict] = []
        self.session_records: List[Dict] = []
        self.frame_samples: List[Dict] = []
        self.alert_records: List[Dict] = []
        self.benchmark_rows: List[Dict] = []
        self.task_queue: List[Dict] = []
        self.queue_running = False
        self.queue_index = 0
        self.class_thresholds: Dict[str, float] = self._load_class_thresholds()
        self.empty_frame_streak = 0
        self.last_box_count_for_alert: Optional[int] = None
        self.record_limit_reached = False
        self.is_running = False
        self.is_paused = False
        self.is_recording = False
        self.theme_name = self.settings.value("ui/theme", "light")

        self.temp_root = ensure_dir(os.path.join(tempfile.gettempdir(), self.config.temp_dir_name))
        self.temp_video_path = os.path.join(self.temp_root, "detected_video.mp4")
        self.temp_camera_record_path = os.path.join(self.temp_root, "camera_record.mp4")
        self.temp_folder_dir = os.path.join(self.temp_root, "folder_results")
        self.temp_image_path = os.path.join(self.temp_root, "detected_image.jpg")
        self.alert_capture_dir = ensure_dir(str(self.project_root / self.config.backend_dir_name / "alert_captures"))
        clear_dir(self.temp_folder_dir)

        self._build_ui()
        self._restore_settings()
        self._apply_theme()
        self._bind_initial_state()
        self.resource_timer = QTimer(self)
        self.resource_timer.timeout.connect(self._update_resource_monitor)
        self.resource_timer.start(self.config.resource_poll_interval_ms)
        self._update_resource_monitor()

    # --------------------------- UI ---------------------------
    def _build_ui(self) -> None:
        self.setWindowTitle(self.config.window_title)
        self.setWindowFlag(Qt.WindowMaximizeButtonHint, False)

        self.content_width = 540
        self.image_gap = 20
        self.image_width = int(self.content_width * 1.2)
        self.image_height = int(self.content_width * 0.94)
        self.card_gap = 10
        self.card_width = (self.content_width - self.card_gap * 2) // 3
        self.left_pair_gap = 20
        self.main_panel_width = self.image_width
        self.selector_button_width = (self.main_panel_width - 3 * 10) // 4
        self.path_label_width = self.main_panel_width - self.selector_button_width - 10
        self.selector_row_height = 36
        self.action_button_height = 30
        self.left_info_panel_width = 205
        self.right_panel_width = 370
        self.column_gap = 2
        self.root_side_margin = 6
        self.right_group_spacing = 6
        self.right_group_inner_margin = 6
        self.right_group_label_height = 26
        self.group_row_gap = self.right_group_spacing
        self.top_section_height = 210
        fit_window_width = (
            self.left_info_panel_width
            + self.main_panel_width
            + self.right_panel_width
            + self.column_gap * 2
            + self.root_side_margin * 2
        )
        self.setFixedSize(fit_window_width, self.config.window_height)

        title_font = QFont("FangSong", 18)
        title_font.setBold(True)
        display_font = QFont("FangSong", 24)
        display_font.setBold(True)

        self.title_label = QLabel("Visual Interface for YOLO")
        self.title_label.setAlignment(Qt.AlignCenter)
        self.title_label.setFixedSize(self.main_panel_width, 50)
        self.title_label.setFont(title_font)

        self.label_original = QLabel("初始图像")
        self.label_original.setAlignment(Qt.AlignCenter)
        self.label_original.setFixedSize(self.image_width, self.image_height)
        self.label_original.setFont(display_font)
        self.label_original.hide()

        self.label_result = QLabel("检测图像")
        self.label_result.setAlignment(Qt.AlignCenter)
        self.label_result.setFixedSize(self.image_width, self.image_height)
        self.label_result.setFont(display_font)

        image_row = QHBoxLayout()
        image_row.setSpacing(8)
        image_row.setContentsMargins(0, 0, 0, 0)
        image_row.addWidget(self.label_original, 0, Qt.AlignCenter)
        image_row.addWidget(self.label_result, 0, Qt.AlignCenter)

        self.box_count_card = QLabel("检测框数量: 0")
        self.box_count_card.setAlignment(Qt.AlignCenter)
        self.box_count_card.setFixedSize(self.card_width, 44)
        self.box_count_card.setFont(QFont("Arial", 12, QFont.Bold))

        self.fps_card = QLabel("FPS: 0.0")
        self.fps_card.setAlignment(Qt.AlignCenter)
        self.fps_card.setFixedSize(self.card_width, 44)
        self.fps_card.setFont(QFont("Arial", 12, QFont.Bold))

        self.mode_card = QLabel("当前模式: 未选择")
        self.mode_card.setAlignment(Qt.AlignCenter)
        self.mode_card.setFixedSize(self.card_width, 44)
        self.mode_card.setFont(QFont("Arial", 12, QFont.Bold))

        card_row = QHBoxLayout()
        card_row.setSpacing(self.card_gap)
        card_row.setContentsMargins(0, 0, 0, 0)
        card_row.addWidget(self.box_count_card)
        card_row.addWidget(self.fps_card)
        card_row.addWidget(self.mode_card)

        self.progress_bar = QProgressBar()
        self.progress_bar.setValue(0)
        self.progress_bar.setFormat("准备就绪")
        self.progress_bar.setFixedHeight(20)
        self.progress_bar.setFixedWidth(self.main_panel_width)

        self.load_model_button = QPushButton("模型选择")
        self.load_model_button.clicked.connect(self.select_model)
        self.model_path_label = self._create_path_label("未选择模型")

        self.image_detect_button = QPushButton("图片检测")
        self.image_detect_button.clicked.connect(self.select_image)
        self.image_path_label = self._create_path_label("未选择图片")

        self.folder_detect_button = QPushButton("文件夹检测")
        self.folder_detect_button.clicked.connect(self.select_folder)
        self.folder_path_label = self._create_path_label("未选择文件夹")

        self.video_detect_button = QPushButton("视频检测")
        self.video_detect_button.clicked.connect(self.select_video)
        self.video_path_label = self._create_path_label("未选择视频")

        self.camera_detect_button = QPushButton("摄像头检测")
        self.camera_detect_button.clicked.connect(self.select_camera)
        self.remote_stream_button = QPushButton("远程流")
        self.remote_stream_button.clicked.connect(self.select_remote_stream)
        self.save_button = QPushButton("保存结果")
        self.save_button.clicked.connect(self.handle_save_action)
        self.export_button = QPushButton("导出日志")
        self.export_button.clicked.connect(self.export_log)
        self.analysis_button = QPushButton("实验分析")
        self.analysis_button.clicked.connect(self.open_analysis_center)
        self.start_button = QPushButton("开始检测")
        self.start_button.clicked.connect(self.start_detection)
        self.pause_button = QPushButton("暂停")
        self.pause_button.clicked.connect(self.toggle_pause)
        self.theme_button = QPushButton("切换主题")
        self.theme_button.clicked.connect(self.toggle_theme)
        self.reset_button = QPushButton("重置界面")
        self.reset_button.clicked.connect(self.reset_view)

        self.load_model_button.setText("模型载入")
        self.image_detect_button.setText("图片检测")
        self.folder_detect_button.setText("文件夹检测")
        self.video_detect_button.setText("视频检测")
        self.camera_detect_button.setText("摄像头检测")
        self.remote_stream_button.setText("远程流")

        selector_grid = self._build_selector_grid()
        action_grid = self._build_action_grid()

        action_frame = QFrame()
        action_frame.setObjectName("actionFrame")
        action_frame.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)

        action_frame_layout = QVBoxLayout(action_frame)
        action_frame_layout.setContentsMargins(4, 4, 4, 4)
        action_frame_layout.setSpacing(4)
        action_frame_layout.addLayout(action_grid, 1)

        main_layout = QVBoxLayout()
        main_layout.setSpacing(8)
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.addLayout(image_row)
        main_layout.addWidget(self.progress_bar, 0, Qt.AlignLeft)
        main_layout.addLayout(selector_grid)
        main_layout.addStretch(1)

        self.state_model_value = QLabel("未加载")
        self.state_source_value = QLabel("无")
        self.state_status_value = QLabel("空闲")
        self.state_records_value = QLabel("0")
        self.state_resource_value = QLabel("CPU --")
        self.state_memory_value = QLabel("--")
        self.state_gpu_value = QLabel("GPU --")
        self.state_vram_value = QLabel("--")

        state_group = QGroupBox()
        state_layout = QVBoxLayout()
        state_layout.setContentsMargins(self.right_group_inner_margin, 8, self.right_group_inner_margin, 8)
        state_layout.setSpacing(4)

        self.state_model_key = QLabel("模型")
        self.state_model_key.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        self.state_status_key = QLabel("状态")
        self.state_status_key.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        self.state_records_key = QLabel("日志")
        self.state_records_key.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        self.state_resource_key = QLabel("资源")
        self.state_resource_key.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        self.state_memory_key = QLabel("内存")
        self.state_memory_key.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        self.state_gpu_key = QLabel("显卡")
        self.state_gpu_key.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        self.state_vram_key = QLabel("显存")
        self.state_vram_key.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)

        self.state_model_value.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        self.state_status_value.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        self.state_records_value.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        self.state_resource_value.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        self.state_memory_value.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        self.state_gpu_value.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        self.state_vram_value.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)

        state_form = QGridLayout()
        state_form.setContentsMargins(0, 0, 0, 0)
        state_form.setHorizontalSpacing(6)
        state_form.setVerticalSpacing(4)
        for row, (key_label, value_label) in enumerate(
            [
                (self.state_model_key, self.state_model_value),
                (self.state_status_key, self.state_status_value),
                (self.state_records_key, self.state_records_value),
                (self.state_resource_key, self.state_resource_value),
                (self.state_memory_key, self.state_memory_value),
                (self.state_gpu_key, self.state_gpu_value),
                (self.state_vram_key, self.state_vram_value),
            ]
        ):
            key_label.setFixedWidth(38)
            state_form.addWidget(key_label, row, 0)
            state_form.addWidget(value_label, row, 1)
        state_form.setColumnStretch(1, 1)

        state_layout.addWidget(self._create_section_label("运行状态"))
        state_layout.addLayout(state_form)
        state_layout.addStretch(1)
        state_group.setLayout(state_layout)
        state_group.setFixedHeight(255)

        params_group = QGroupBox()
        self.conf_spin = QDoubleSpinBox()
        self.conf_spin.setRange(0.01, 1.00)
        self.conf_spin.setSingleStep(0.01)
        self.conf_spin.setFixedHeight(30)
        self.iou_spin = QDoubleSpinBox()
        self.iou_spin.setRange(0.01, 1.00)
        self.iou_spin.setSingleStep(0.01)
        self.iou_spin.setFixedHeight(30)
        self.imgsz_spin = QSpinBox()
        self.imgsz_spin.setRange(160, 1280)
        self.imgsz_spin.setSingleStep(32)
        self.imgsz_spin.setValue(self.config.inference_imgsz)
        self.imgsz_spin.setFixedHeight(30)
        self.show_labels_checkbox = QCheckBox("显示类别标签")
        self.show_conf_checkbox = QCheckBox("显示置信度")
        self.compare_view_checkbox = QCheckBox("原图对比")
        self.compare_view_checkbox.toggled.connect(self._toggle_compare_view)
        self.heatmap_checkbox = QCheckBox("检测热力图")
        self.save_folder_results_checkbox = QCheckBox("全量保存(慢)")
        self.roi_alert_checkbox = QCheckBox("显示ROI/报警")
        self.roi_alert_checkbox.toggled.connect(self._refresh_roi_preview)
        self.alert_mark_checkbox = QCheckBox("异常标记")
        self.auto_capture_checkbox = QCheckBox("异常自动截图")
        self.alarm_rules_checkbox = QCheckBox("规则报警")
        self.low_conf_spin = QDoubleSpinBox()
        self.low_conf_spin.setRange(0.01, 1.00)
        self.low_conf_spin.setSingleStep(0.01)
        self.low_conf_spin.setValue(self.config.default_low_confidence_alert)
        self.low_conf_spin.setFixedHeight(26)
        self.box_jump_spin = QSpinBox()
        self.box_jump_spin.setRange(1, 99)
        self.box_jump_spin.setValue(self.config.default_box_jump_alert)
        self.box_jump_spin.setFixedHeight(26)
        self.no_target_spin = QSpinBox()
        self.no_target_spin.setRange(1, 999)
        self.no_target_spin.setValue(self.config.default_no_target_alert_frames)
        self.no_target_spin.setFixedHeight(26)
        self.class_count_spin = QSpinBox()
        self.class_count_spin.setRange(1, 999)
        self.class_count_spin.setValue(self.config.default_class_count_alert)
        self.class_count_spin.setFixedHeight(26)
        roi_x, roi_y, roi_w, roi_h = self.config.default_roi_percent
        self.roi_x_spin = QSpinBox()
        self.roi_y_spin = QSpinBox()
        self.roi_w_spin = QSpinBox()
        self.roi_h_spin = QSpinBox()
        for spin, value in [
            (self.roi_x_spin, roi_x),
            (self.roi_y_spin, roi_y),
            (self.roi_w_spin, roi_w),
            (self.roi_h_spin, roi_h),
        ]:
            spin.setRange(0, 100)
            spin.setValue(value)
            spin.setFixedHeight(24)
            spin.valueChanged.connect(self._refresh_roi_preview)
        self.show_labels_checkbox.setChecked(True)
        self.show_conf_checkbox.setChecked(True)
        self.heatmap_checkbox.setChecked(False)
        self.save_folder_results_checkbox.setChecked(self.config.save_folder_results)
        self.roi_alert_checkbox.setChecked(False)
        self.alert_mark_checkbox.setChecked(True)
        self.auto_capture_checkbox.setChecked(True)
        self.alarm_rules_checkbox.setChecked(False)
        self.show_labels_checkbox.setFixedHeight(24)
        self.show_conf_checkbox.setFixedHeight(24)
        self.compare_view_checkbox.setFixedHeight(24)
        self.heatmap_checkbox.setFixedHeight(24)
        self.save_folder_results_checkbox.setFixedHeight(24)
        self.roi_alert_checkbox.setFixedHeight(24)
        self.alert_mark_checkbox.setFixedHeight(24)
        self.auto_capture_checkbox.setFixedHeight(24)
        self.alarm_rules_checkbox.setFixedHeight(24)
        params_layout = QVBoxLayout()
        params_layout.setContentsMargins(self.right_group_inner_margin, 8, self.right_group_inner_margin, 10)
        params_layout.setSpacing(5)

        params_layout.addWidget(self._create_section_label("参数设置"))

        params_tabs = QTabWidget()
        params_tabs.setFixedHeight(236)

        basic_tab = QWidget()
        basic_layout = QVBoxLayout(basic_tab)
        basic_layout.setContentsMargins(4, 4, 4, 4)
        basic_layout.setSpacing(4)
        basic_grid = QGridLayout()
        basic_grid.setContentsMargins(0, 0, 0, 0)
        basic_grid.setHorizontalSpacing(5)
        basic_grid.setVerticalSpacing(3)
        basic_grid.addWidget(QLabel("置信度"), 0, 0)
        basic_grid.addWidget(self.conf_spin, 0, 1)
        basic_grid.addWidget(QLabel("IoU"), 1, 0)
        basic_grid.addWidget(self.iou_spin, 1, 1)
        basic_grid.addWidget(QLabel("尺寸"), 2, 0)
        basic_grid.addWidget(self.imgsz_spin, 2, 1)
        basic_layout.addLayout(basic_grid)
        basic_layout.addWidget(self.show_labels_checkbox)
        basic_layout.addWidget(self.show_conf_checkbox)
        basic_layout.addWidget(self.compare_view_checkbox)
        basic_layout.addWidget(self.heatmap_checkbox)
        basic_layout.addStretch(1)

        extended_tab = QWidget()
        extended_layout = QVBoxLayout(extended_tab)
        extended_layout.setContentsMargins(4, 4, 4, 4)
        extended_layout.setSpacing(4)
        extended_layout.addWidget(self.save_folder_results_checkbox)
        extended_layout.addWidget(self.roi_alert_checkbox)
        extended_layout.addWidget(self.auto_capture_checkbox)
        roi_grid = QGridLayout()
        roi_grid.setContentsMargins(0, 0, 0, 0)
        roi_grid.setHorizontalSpacing(4)
        roi_grid.setVerticalSpacing(2)
        for col, (label_text, spin) in enumerate(
            [("X%", self.roi_x_spin), ("Y%", self.roi_y_spin), ("W%", self.roi_w_spin), ("H%", self.roi_h_spin)]
        ):
            roi_grid.addWidget(QLabel(label_text), 0, col)
            roi_grid.addWidget(spin, 1, col)
        extended_layout.addLayout(roi_grid)
        extended_layout.addStretch(1)

        alarm_tab = QWidget()
        alarm_layout = QVBoxLayout(alarm_tab)
        alarm_layout.setContentsMargins(4, 4, 4, 4)
        alarm_layout.setSpacing(4)
        alarm_layout.addWidget(self.alert_mark_checkbox)
        alarm_layout.addWidget(self.alarm_rules_checkbox)
        alert_grid = QGridLayout()
        alert_grid.setContentsMargins(0, 0, 0, 0)
        alert_grid.setHorizontalSpacing(8)
        alert_grid.setVerticalSpacing(4)
        alert_grid.addWidget(QLabel("低置信"), 0, 0)
        alert_grid.addWidget(self.low_conf_spin, 0, 1)
        alert_grid.addWidget(QLabel("数量突变"), 1, 0)
        alert_grid.addWidget(self.box_jump_spin, 1, 1)
        alert_grid.addWidget(QLabel("无目标帧"), 2, 0)
        alert_grid.addWidget(self.no_target_spin, 2, 1)
        alert_grid.addWidget(QLabel("类别数量"), 3, 0)
        alert_grid.addWidget(self.class_count_spin, 3, 1)
        alert_grid.setColumnStretch(1, 1)
        alarm_layout.addLayout(alert_grid)
        alarm_layout.addStretch(1)

        params_tabs.addTab(basic_tab, "检测")
        params_tabs.addTab(extended_tab, "扩展")
        params_tabs.addTab(alarm_tab, "报警")
        params_layout.addWidget(params_tabs)
        params_layout.addStretch(1)
        params_group.setLayout(params_layout)
        params_group.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)

        stats_group = QGroupBox()
        self.summary_label = QLabel("当前暂无检测结果")
        self.summary_label.setWordWrap(True)
        self.summary_label.setAlignment(Qt.AlignTop | Qt.AlignLeft)
        self.summary_label.setMinimumHeight(92)
        summary_layout = QVBoxLayout()
        summary_layout.setContentsMargins(self.right_group_inner_margin, 10, self.right_group_inner_margin, 12)
        summary_layout.addWidget(self._create_section_label("结果摘要"))
        summary_layout.addWidget(self.summary_label)
        stats_group.setLayout(summary_layout)
        stats_group.setFixedHeight(self.top_section_height)

        fps_group = QGroupBox()
        self.summary_fps_label = QLabel("FPS: 0.0")
        self.summary_fps_label.setAlignment(Qt.AlignCenter)
        self.summary_fps_label.setFixedHeight(34)
        self.summary_fps_label.setFont(QFont("Arial", 13, QFont.Bold))
        fps_layout = QVBoxLayout()
        fps_layout.setContentsMargins(self.right_group_inner_margin, 8, self.right_group_inner_margin, 8)
        fps_layout.setSpacing(4)
        fps_layout.addWidget(self._create_section_label("FPS"))
        fps_layout.addWidget(self.summary_fps_label)
        fps_group.setLayout(fps_layout)
        fps_group.setFixedHeight(78)

        table_group = QGroupBox()
        self.result_table = QTableWidget(0, 3)
        self.result_table.setHorizontalHeaderLabels(["序号", "类别", "置信度"])
        self.result_table.verticalHeader().setVisible(False)
        self.result_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.result_table.setSelectionBehavior(QTableWidget.SelectRows)
        self.result_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeToContents)
        self.result_table.horizontalHeader().setSectionResizeMode(1, QHeaderView.Stretch)
        self.result_table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeToContents)
        self.result_table.horizontalHeader().setFixedHeight(32)
        self.result_table.verticalHeader().setDefaultSectionSize(28)
        table_layout = QVBoxLayout()
        table_layout.setContentsMargins(self.right_group_inner_margin, 10, self.right_group_inner_margin, 12)
        table_layout.addWidget(self._create_section_label("检测明细"))
        table_layout.addWidget(self.result_table)
        table_group.setLayout(table_layout)
        table_group.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)

        params_group.setFixedHeight(285)

        left_info_layout = QVBoxLayout()
        left_info_layout.setSpacing(self.group_row_gap)
        left_info_layout.setContentsMargins(0, 0, 0, 0)
        left_info_layout.addWidget(stats_group, 0)
        left_info_layout.addWidget(fps_group, 0)
        left_info_layout.addWidget(table_group, 1)

        right_layout = QVBoxLayout()
        right_layout.setSpacing(self.group_row_gap)
        right_layout.setContentsMargins(0, 0, 0, 0)
        right_layout.addWidget(params_group, 0)
        right_layout.addWidget(state_group, 0)
        right_layout.addWidget(action_frame, 1)

        left_info_widget = QWidget()
        left_info_widget.setFixedWidth(self.left_info_panel_width)
        left_info_widget.setSizePolicy(QSizePolicy.Fixed, QSizePolicy.Expanding)
        left_info_widget.setLayout(left_info_layout)

        main_widget = QWidget()
        main_widget.setFixedWidth(self.main_panel_width)
        main_widget.setSizePolicy(QSizePolicy.Fixed, QSizePolicy.Expanding)
        main_widget.setLayout(main_layout)

        right_widget = QWidget()
        right_widget.setFixedWidth(self.right_panel_width)
        right_widget.setSizePolicy(QSizePolicy.Fixed, QSizePolicy.Expanding)
        right_widget.setLayout(right_layout)

        main_row = QHBoxLayout()
        main_row.setSpacing(self.column_gap)
        main_row.setContentsMargins(0, 0, 0, 0)
        main_row.addWidget(main_widget)
        main_row.addWidget(left_info_widget)
        main_row.addWidget(right_widget)

        root_layout = QVBoxLayout()
        root_layout.setSpacing(0)
        root_layout.setContentsMargins(self.root_side_margin, 0, self.root_side_margin, 0)
        root_layout.addLayout(main_row, 1)

        central = QWidget()
        central.setObjectName("centralWidget")
        central.setLayout(root_layout)
        self.setCentralWidget(central)
        self.statusBar().showMessage("请选择模型")

    def _create_section_label(self, text: str) -> QLabel:
        label = QLabel(text)
        label.setObjectName("sectionLabel")
        label.setAlignment(Qt.AlignCenter)
        label.setFixedHeight(24)
        label.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        return label

    def _create_path_label(self, text: str) -> QLabel:
        label = QLabel(text)
        label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        label.setWordWrap(False)
        label.setFixedSize(self.path_label_width, self.selector_row_height)
        label.setSizePolicy(QSizePolicy.Fixed, QSizePolicy.Fixed)
        return label

    def _build_selector_grid(self) -> QVBoxLayout:
        wrapper = QVBoxLayout()
        wrapper.setSpacing(4)
        wrapper.setContentsMargins(0, 0, 0, 0)

        rows = [
            (self.load_model_button, self.model_path_label),
            (self.image_detect_button, self.image_path_label),
            (self.folder_detect_button, self.folder_path_label),
            (self.video_detect_button, self.video_path_label),
        ]
        for button, label in rows:
            row = QHBoxLayout()
            row.setSpacing(6)
            row.setContentsMargins(0, 0, 0, 0)
            row.addWidget(button)
            row.addWidget(label)
            wrapper.addLayout(row)

        return wrapper

    def _build_action_grid(self) -> QGridLayout:
        grid = QGridLayout()
        grid.setHorizontalSpacing(6)
        grid.setVerticalSpacing(6)
        grid.setContentsMargins(0, 0, 0, 0)

        buttons = [
            self.camera_detect_button,
            self.remote_stream_button,
            self.start_button,
            self.pause_button,
            self.save_button,
            self.export_button,
            self.analysis_button,
            self.reset_button,
            self.theme_button,
        ]
        for index, button in enumerate(buttons):
            row = index // 3
            col = index % 3
            grid.addWidget(button, row, col)
        grid.setRowStretch(0, 1)
        grid.setRowStretch(1, 1)
        grid.setRowStretch(2, 1)
        grid.setColumnStretch(0, 1)
        grid.setColumnStretch(1, 1)
        grid.setColumnStretch(2, 1)

        return grid

    # ----------------------- Theme and settings -----------------------
    def _apply_theme(self) -> None:
        theme = theme_palette(self.theme_name)
        self.setStyleSheet(main_window_styles(theme))

        self.title_label.setStyleSheet(card_style(theme))
        self.label_original.setStyleSheet(card_style(theme))
        self.label_result.setStyleSheet(card_style(theme))
        self.box_count_card.setStyleSheet(card_style(theme))
        self.fps_card.setStyleSheet(card_style(theme))
        self.mode_card.setStyleSheet(card_style(theme))
        self.model_path_label.setStyleSheet(path_label_style(theme))
        self.image_path_label.setStyleSheet(path_label_style(theme))
        self.folder_path_label.setStyleSheet(path_label_style(theme))
        self.video_path_label.setStyleSheet(path_label_style(theme))

        model_label_font = self.model_path_label.font()
        path_font = QFont(model_label_font)
        path_font.setPointSize(9)
        compact_font = QFont(model_label_font)
        compact_font.setPointSize(9)
        action_font = QFont(model_label_font)
        action_font.setPointSize(9)
        for path_label in [
            self.model_path_label,
            self.image_path_label,
            self.folder_path_label,
            self.video_path_label,
        ]:
            path_label.setFont(path_font)

        for key_label in [
            self.state_model_key,
            self.state_status_key,
            self.state_records_key,
            self.state_resource_key,
            self.state_memory_key,
            self.state_gpu_key,
            self.state_vram_key,
        ]:
            key_label.setFont(compact_font)

        btns = [
            (self.load_model_button, None),
            (self.image_detect_button, None),
            (self.folder_detect_button, None),
            (self.video_detect_button, None),
            (self.camera_detect_button, None),
            (self.remote_stream_button, None),
            (self.save_button, None),
            (self.export_button, None),
            (self.analysis_button, None),
            (self.pause_button, None),
            (self.theme_button, None),
            (self.start_button, None),
            (self.reset_button, None),
        ]
        param_widgets = [
            self.conf_spin,
            self.iou_spin,
            self.imgsz_spin,
            self.show_labels_checkbox,
            self.show_conf_checkbox,
            self.compare_view_checkbox,
            self.heatmap_checkbox,
            self.save_folder_results_checkbox,
            self.roi_alert_checkbox,
            self.alert_mark_checkbox,
            self.auto_capture_checkbox,
            self.alarm_rules_checkbox,
            self.low_conf_spin,
            self.box_jump_spin,
            self.no_target_spin,
            self.class_count_spin,
            self.roi_x_spin,
            self.roi_y_spin,
            self.roi_w_spin,
            self.roi_h_spin,
        ]
        for widget in param_widgets:
            widget.setFont(compact_font)
        selector_buttons = {
            self.load_model_button,
            self.image_detect_button,
            self.folder_detect_button,
            self.video_detect_button,
        }
        action_buttons = {
            self.camera_detect_button,
            self.remote_stream_button,
            self.save_button,
            self.export_button,
            self.analysis_button,
            self.start_button,
            self.pause_button,
            self.theme_button,
            self.reset_button,
        }
        for button, bg in btns:
            button.setFont(action_font)
            button.setStyleSheet(button_style(theme, background_color=bg))
            if button in selector_buttons:
                button.setFixedSize(self.selector_button_width, self.selector_row_height)
            elif button in action_buttons:
                button.setMinimumHeight(self.action_button_height)
                button.setMaximumHeight(16777215)
                button.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)

        for value_label in [
            self.state_model_value,
            self.state_source_value,
            self.state_status_value,
            self.state_records_value,
            self.state_resource_value,
            self.state_memory_value,
            self.state_gpu_value,
            self.state_vram_value,
            self.summary_fps_label,
        ]:
            if value_label is self.state_model_value:
                value_label.setFixedHeight(self.right_group_label_height + 8)
            else:
                value_label.setFixedHeight(self.right_group_label_height)
            if value_label is self.summary_fps_label:
                value_label.setAlignment(Qt.AlignCenter)
            else:
                value_label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
            value_label.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
            value_label.setFont(compact_font)
            value_label.setStyleSheet(
                    f"""
                    background: {theme['status_bg']};
                    border: 1px solid {theme['soft_border']};
                    color: {theme['text_fg']};
                    padding: 2px 5px;
                    font-weight: 500;
                    """
                )
        self.summary_label.setStyleSheet(f"padding: 2px 2px; line-height: 1.35; color: {theme['text_fg']};")

    def toggle_theme(self) -> None:
        self.theme_name = "dark" if self.theme_name == "light" else "light"
        self._apply_theme()
        self.settings.setValue("ui/theme", self.theme_name)

    def _restore_settings(self) -> None:
        self.conf_spin.setValue(float(self.settings.value("params/conf", self.config.default_confidence)))
        self.iou_spin.setValue(float(self.settings.value("params/iou", self.config.default_iou)))
        self.imgsz_spin.setValue(int(self.settings.value("params/imgsz", self.config.inference_imgsz)))
        self.show_labels_checkbox.setChecked(self.settings.value("params/show_labels", True, type=bool))
        self.show_conf_checkbox.setChecked(self.settings.value("params/show_conf", True, type=bool))
        self.compare_view_checkbox.setChecked(self.settings.value("params/compare_view", False, type=bool))
        self.heatmap_checkbox.setChecked(self.settings.value("params/heatmap", False, type=bool))
        self.save_folder_results_checkbox.setChecked(False)
        self.roi_alert_checkbox.setChecked(self.settings.value("params/roi_alert", False, type=bool))
        self.alert_mark_checkbox.setChecked(self.settings.value("params/alert_mark", True, type=bool))
        self.auto_capture_checkbox.setChecked(self.settings.value("params/auto_capture", True, type=bool))
        self.alarm_rules_checkbox.setChecked(self.settings.value("params/alarm_rules", False, type=bool))
        self.low_conf_spin.setValue(float(self.settings.value("params/low_conf", self.config.default_low_confidence_alert)))
        self.box_jump_spin.setValue(int(self.settings.value("params/box_jump", self.config.default_box_jump_alert)))
        self.no_target_spin.setValue(int(self.settings.value("params/no_target_frames", self.config.default_no_target_alert_frames)))
        self.class_count_spin.setValue(int(self.settings.value("params/class_count_alert", self.config.default_class_count_alert)))
        roi_x, roi_y, roi_w, roi_h = self.config.default_roi_percent
        self.roi_x_spin.setValue(int(self.settings.value("params/roi_x", roi_x)))
        self.roi_y_spin.setValue(int(self.settings.value("params/roi_y", roi_y)))
        self.roi_w_spin.setValue(int(self.settings.value("params/roi_w", roi_w)))
        self.roi_h_spin.setValue(int(self.settings.value("params/roi_h", roi_h)))

    def _save_settings(self) -> None:
        self.settings.setValue("params/conf", self.conf_spin.value())
        self.settings.setValue("params/iou", self.iou_spin.value())
        self.settings.setValue("params/imgsz", self.imgsz_spin.value())
        self.settings.setValue("params/show_labels", self.show_labels_checkbox.isChecked())
        self.settings.setValue("params/show_conf", self.show_conf_checkbox.isChecked())
        self.settings.setValue("params/compare_view", self.compare_view_checkbox.isChecked())
        self.settings.setValue("params/heatmap", self.heatmap_checkbox.isChecked())
        self.settings.setValue("params/save_folder_results", False)
        self.settings.setValue("params/roi_alert", self.roi_alert_checkbox.isChecked())
        self.settings.setValue("params/alert_mark", self.alert_mark_checkbox.isChecked())
        self.settings.setValue("params/auto_capture", self.auto_capture_checkbox.isChecked())
        self.settings.setValue("params/alarm_rules", self.alarm_rules_checkbox.isChecked())
        self.settings.setValue("params/low_conf", self.low_conf_spin.value())
        self.settings.setValue("params/box_jump", self.box_jump_spin.value())
        self.settings.setValue("params/no_target_frames", self.no_target_spin.value())
        self.settings.setValue("params/class_count_alert", self.class_count_spin.value())
        self.settings.setValue("params/roi_x", self.roi_x_spin.value())
        self.settings.setValue("params/roi_y", self.roi_y_spin.value())
        self.settings.setValue("params/roi_w", self.roi_w_spin.value())
        self.settings.setValue("params/roi_h", self.roi_h_spin.value())
        self.settings.setValue("ui/theme", self.theme_name)

    def _bind_initial_state(self) -> None:
        for widget in [
            self.image_detect_button,
            self.folder_detect_button,
            self.video_detect_button,
            self.camera_detect_button,
            self.remote_stream_button,
            self.start_button,
            self.pause_button,
            self.save_button,
            self.export_button,
        ]:
            widget.setEnabled(False)

    # ----------------------- Selection -----------------------
    def select_model(self) -> None:
        file_path, _ = QFileDialog.getOpenFileName(self, "选择模型文件", "", "权重文件 (*.pt *.onnx *.engine *.torchscript)")
        if not file_path:
            return
        try:
            self.detector.load_model(file_path)
        except Exception as exc:
            QMessageBox.critical(self, "加载失败", f"模型加载失败：\n{exc}")
            return

        self._update_path_label(self.model_path_label, self.detector.model_path or file_path)
        name = Path(file_path).name
        self.state_model_value.setText(f"{name} ({self.detector.runtime_label()}, {self._current_imgsz()})")
        self.statusBar().showMessage(f"模型已加载：{name}，后端：{self.detector.runtime_label()}，输入尺寸：{self._current_imgsz()}")
        for widget in [
            self.image_detect_button,
            self.folder_detect_button,
            self.video_detect_button,
            self.camera_detect_button,
            self.remote_stream_button,
        ]:
            widget.setEnabled(True)

    def select_image(self) -> None:
        path, _ = QFileDialog.getOpenFileName(self, "选择图片", "", "图片文件 (*.jpg *.jpeg *.png *.bmp *.webp *.tif *.tiff)")
        if not path:
            return
        self._set_image_source(path)

    def _set_image_source(self, path: str) -> None:
        self._stop_worker_if_needed()
        self.current_mode = DetectionMode.IMAGE
        self.current_image_path = path
        self.current_folder_images = []
        self.current_video_path = None
        self.current_camera_id = None
        self._update_path_label(self.image_path_label, path)
        self._clear_other_paths("image")
        image = cv2.imread(path)
        if image is not None:
            self.last_original_image = image
            self._refresh_image_panels(image, None)
            self.label_result.setText("检测图像")
        self.mode_card.setText("当前模式: 图片")
        self.state_source_value.setText(Path(path).name)
        self.state_status_value.setText("待检测")
        self.start_button.setEnabled(True)
        self.pause_button.setEnabled(False)
        self.save_button.setEnabled(False)
        self.export_button.setEnabled(False)

    def select_folder(self) -> None:
        folder = QFileDialog.getExistingDirectory(self, "选择图片文件夹")
        if not folder:
            return
        self._set_folder_source(folder)

    def _set_folder_source(self, folder: str) -> None:
        images = list_images(folder, self.config.supported_image_exts)
        if not images:
            QMessageBox.warning(self, "提示", "该文件夹中没有可检测的图片。")
            return
        self._stop_worker_if_needed()
        self.current_mode = DetectionMode.FOLDER
        self.current_folder_images = images
        self.current_image_path = None
        self.current_video_path = None
        self.current_camera_id = None
        self._update_path_label(self.folder_path_label, folder)
        self._clear_other_paths("folder")
        first = cv2.imread(images[0])
        if first is not None:
            self.last_original_image = first
            self._refresh_image_panels(first, None)
        self.mode_card.setText("当前模式: 文件夹")
        self.state_source_value.setText(Path(folder).name)
        self.state_status_value.setText(f"待检测（{len(images)} 张）")
        self.progress_bar.setValue(0)
        self.progress_bar.setFormat(f"待处理 0 / {len(images)}")
        self.start_button.setEnabled(True)
        self.pause_button.setEnabled(False)
        self.save_button.setEnabled(False)
        self.export_button.setEnabled(False)

    def select_video(self) -> None:
        path, _ = QFileDialog.getOpenFileName(self, "选择视频", "", "视频文件 (*.mp4 *.avi *.mov *.mkv *.flv *.wmv)")
        if not path:
            return
        self._set_video_source(path, remote=False)

    def _set_video_source(self, path: str, remote: bool = False) -> None:
        self._stop_worker_if_needed()
        self.current_mode = DetectionMode.VIDEO
        self.current_video_path = path
        self.current_folder_images = []
        self.current_image_path = None
        self.current_camera_id = None
        self._update_path_label(self.video_path_label, path)
        self._clear_other_paths("video")
        self.mode_card.setText("当前模式: 远程流" if remote else "当前模式: 视频")
        self.state_source_value.setText(path if remote else Path(path).name)
        self.state_status_value.setText("待检测")
        self.start_button.setEnabled(True)
        self.pause_button.setEnabled(False)
        self.save_button.setEnabled(False)
        self.export_button.setEnabled(False)

    def select_remote_stream(self) -> None:
        url, ok = QInputDialog.getText(
            self,
            "远程视频流",
            "请输入连续视频流地址，例如：\n"
            "rtsp://user:pass@192.168.1.64:554/xxx\n"
            "http://192.168.1.10:8080/video\n\n"
            "注意：不要填写监控网页地址或单张截图地址。",
        )
        if not ok or not url.strip():
            return
        self._set_video_source(url.strip(), remote=True)
        self.statusBar().showMessage("已选择远程视频流，点击开始检测后将尝试连接")

    def select_camera(self) -> None:
        cameras = self._detect_local_cameras()
        if not cameras:
            QMessageBox.warning(self, "提示", "未检测到可用摄像头。")
            return
        camera_texts = [str(c) for c in cameras]
        camera_id_str, ok = QInputDialog.getItem(self, "选择摄像头", "请选择可用摄像头：", camera_texts, 0, False)
        if not ok:
            return
        self._stop_worker_if_needed()
        self.current_mode = DetectionMode.CAMERA
        self.current_camera_id = int(camera_id_str)
        self.current_folder_images = []
        self.current_image_path = None
        self.current_video_path = None
        self._clear_other_paths("camera")
        self.mode_card.setText(f"当前模式: 摄像头 #{camera_id_str}")
        self.state_source_value.setText(f"Camera {camera_id_str}")
        self.state_status_value.setText("待检测")
        self.start_button.setEnabled(True)
        self.pause_button.setEnabled(False)
        self.save_button.setEnabled(False)
        self.export_button.setEnabled(False)
        self.statusBar().showMessage(f"已选择摄像头 {camera_id_str}")

    # ----------------------- Detection -----------------------
    def start_detection(self) -> None:
        if not self.detector.is_loaded:
            QMessageBox.warning(self, "提示", "请先加载模型。")
            return
        if self.current_mode == DetectionMode.NONE:
            QMessageBox.warning(self, "提示", "请先选择图片、文件夹、视频或摄像头。")
            return
        if self.current_mode == DetectionMode.FOLDER and self.save_folder_results_checkbox.isChecked():
            QMessageBox.information(
                self,
                "全量保存提示",
                "全量保存会为每张图片绘制并写入结果图，FPS 会明显下降。需要最高速度时请取消勾选。",
            )

        self._save_settings()
        self._reset_session_records()
        self._set_fps_value(0.0)
        self.box_count_card.setText("检测框数量: 0")
        self.summary_label.setText("正在准备检测任务…")
        self.result_table.setRowCount(0)
        self.state_status_value.setText("运行中")
        self.is_running = True
        self.is_paused = False
        self.pause_button.setText("暂停")
        self._start_backend_session()

        if self.current_mode == DetectionMode.IMAGE:
            self._run_single_image_detection()
            return

        clear_dir(self.temp_folder_dir)
        if os.path.exists(self.temp_video_path):
            os.remove(self.temp_video_path)
        if os.path.exists(self.temp_camera_record_path):
            os.remove(self.temp_camera_record_path)

        source = {
            DetectionMode.FOLDER: self.current_folder_images,
            DetectionMode.VIDEO: self.current_video_path,
            DetectionMode.CAMERA: self.current_camera_id,
        }[self.current_mode]

        self.worker = DetectionWorker(
            detector=self.detector,
            mode=self.current_mode,
            source=source,
            conf=self.conf_spin.value(),
            iou=self.iou_spin.value(),
            imgsz=self._current_imgsz(),
            show_labels=self.show_labels_checkbox.isChecked(),
            show_conf=self.show_conf_checkbox.isChecked(),
            display_fps=self.config.display_fps,
            stream_emit_interval=self.config.stream_ui_emit_interval,
            folder_emit_interval=self.config.folder_ui_emit_interval,
            write_video_result=self.config.auto_write_video_result,
            save_folder_results=self.save_folder_results_checkbox.isChecked(),
            temp_video_path=self.temp_video_path,
            temp_folder_output_dir=self.temp_folder_dir,
        )
        self.worker.frame_ready.connect(self._handle_frame_payload)
        self.worker.status_changed.connect(self._set_status)
        self.worker.error_occurred.connect(self._handle_worker_error)
        self.worker.detection_finished.connect(self._handle_worker_finished)
        self.worker.recording_ready.connect(self._handle_recording_ready)
        self.worker.start()

        self.start_button.setEnabled(False)
        self.pause_button.setEnabled(True)
        self.save_button.setEnabled(self.current_mode == DetectionMode.CAMERA)
        self.export_button.setEnabled(True)
        if self.current_mode == DetectionMode.CAMERA:
            self.save_button.setText("开始录制")
        else:
            self.save_button.setText("保存结果")

    def _run_single_image_detection(self) -> None:
        try:
            image = cv2.imread(self.current_image_path)
            if image is None:
                raise RuntimeError("无法读取当前图片。")
            import time
            infer_start = time.perf_counter()
            result = self.detector.predict(
                image,
                conf=self.conf_spin.value(),
                iou=self.iou_spin.value(),
                imgsz=self._current_imgsz(),
            )
            annotated = self.detector.render(
                result,
                show_labels=self.show_labels_checkbox.isChecked(),
                show_conf=self.show_conf_checkbox.isChecked(),
            )
            elapsed = max(time.perf_counter() - infer_start, 1e-6)
            rows = self.detector.parse_rows(result)
            payload = {
                "original": image,
                "annotated": annotated,
                "rows": rows,
                "box_count": self.detector.get_box_count(result),
                "info_text": self.detector.build_info_text(rows),
                "source_name": Path(self.current_image_path).name,
                "frame_index": 1,
                "fps": 1.0 / elapsed,
                "progress_current": 1,
                "progress_total": 1,
            }
            self._apply_postprocess_to_payload(payload)
            cv2.imwrite(self.temp_image_path, payload["annotated"])
            self._handle_frame_payload(payload)
            self._handle_worker_finished("图片检测完成")
        except Exception as exc:
            self._handle_worker_error(str(exc))

    def toggle_pause(self) -> None:
        if not self.worker or not self.is_running:
            return
        self.is_paused = not self.is_paused
        if self.is_paused:
            self.worker.pause()
            self.pause_button.setText("继续")
            self.state_status_value.setText("已暂停")
            self.statusBar().showMessage("检测已暂停")
        else:
            self.worker.resume()
            self.pause_button.setText("暂停")
            self.state_status_value.setText("运行中")
            self.statusBar().showMessage("检测已恢复")

    def _toggle_compare_view(self) -> None:
        self._update_compare_layout()
        self._refresh_image_panels(self.last_original_image, self.last_annotated_image)

    def _update_compare_layout(self) -> None:
        if self.compare_view_checkbox.isChecked():
            width = max(240, (self.main_panel_width - 8) // 2)
            self.label_original.show()
            self.label_original.setFixedSize(width, self.image_height)
            self.label_result.setFixedSize(width, self.image_height)
        else:
            self.label_original.hide()
            self.label_result.setFixedSize(self.image_width, self.image_height)

    def _refresh_image_panels(self, original=None, annotated=None) -> None:
        self._update_compare_layout()
        if original is not None and self.compare_view_checkbox.isChecked():
            self.label_original.setPixmap(cv_to_pixmap(original, self.label_original.size()))
        if annotated is not None:
            self.label_result.setPixmap(cv_to_pixmap(annotated, self.label_result.size()))

    def _roi_rect_for_image(self, image):
        if image is None:
            return None
        h, w = image.shape[:2]
        x = int(w * self.roi_x_spin.value() / 100)
        y = int(h * self.roi_y_spin.value() / 100)
        rw = int(w * self.roi_w_spin.value() / 100)
        rh = int(h * self.roi_h_spin.value() / 100)
        x = max(0, min(x, max(0, w - 1)))
        y = max(0, min(y, max(0, h - 1)))
        rw = max(1, min(rw, w - x))
        rh = max(1, min(rh, h - y))
        return x, y, x + rw, y + rh

    def _filter_rows_by_class_threshold(self, rows: List[Dict]) -> List[Dict]:
        if not self.class_thresholds:
            return rows
        filtered = []
        for row in rows:
            class_name = str(row.get("class_name", ""))
            threshold = self.class_thresholds.get(class_name)
            if threshold is None:
                filtered.append(row)
                continue
            if float(row.get("confidence", 0.0) or 0.0) >= float(threshold):
                filtered.append(row)
        return filtered

    def _draw_rows_on_image(self, image, rows: List[Dict]):
        output = image.copy()
        for row in rows:
            bbox = row.get("bbox") or []
            if len(bbox) != 4:
                continue
            x1, y1, x2, y2 = [int(float(value)) for value in bbox]
            cv2.rectangle(output, (x1, y1), (x2, y2), (0, 220, 255), 2)
            label = f"{row.get('class_name', '')} {float(row.get('confidence', 0.0)):.2f}"
            cv2.rectangle(output, (x1, max(0, y1 - 24)), (min(output.shape[1] - 1, x1 + 190), y1), (0, 180, 220), -1)
            cv2.putText(output, label, (x1 + 4, max(14, y1 - 6)), cv2.FONT_HERSHEY_SIMPLEX, 0.5, (0, 0, 0), 1, cv2.LINE_AA)
        return output

    def _apply_detection_heatmap(self, image, rows: List[Dict]):
        if image is None or not rows:
            return image
        height, width = image.shape[:2]
        heat_field = np.zeros((height, width), dtype=np.float32)

        for row in rows:
            bbox = row.get("bbox") or []
            if len(bbox) != 4:
                continue
            x1, y1, x2, y2 = [float(value) for value in bbox]
            x1 = max(0.0, min(float(width - 1), x1))
            x2 = max(0.0, min(float(width - 1), x2))
            y1 = max(0.0, min(float(height - 1), y1))
            y2 = max(0.0, min(float(height - 1), y2))
            box_w = max(1.0, x2 - x1)
            box_h = max(1.0, y2 - y1)
            center = row.get("center") or [(x1 + x2) / 2, (y1 + y2) / 2]
            cx, cy = float(center[0]), float(center[1])
            confidence = float(row.get("confidence", 0.0) or 0.0)
            sigma_x = max(10.0, box_w * 0.85)
            sigma_y = max(10.0, box_h * 0.85)
            margin = int(max(sigma_x, sigma_y) * 2.8)
            px1 = max(0, int(cx - margin))
            px2 = min(width, int(cx + margin + 1))
            py1 = max(0, int(cy - margin))
            py2 = min(height, int(cy + margin + 1))
            if px1 >= px2 or py1 >= py2:
                continue

            xs = np.arange(px1, px2, dtype=np.float32) - cx
            ys = np.arange(py1, py2, dtype=np.float32) - cy
            xx, yy = np.meshgrid(xs, ys)
            patch = np.exp(-0.5 * ((xx / sigma_x) ** 2 + (yy / sigma_y) ** 2)) * max(0.15, confidence)
            heat_field[py1:py2, px1:px2] += patch.astype(np.float32)

        if float(heat_field.max()) <= 0:
            return image

        heat_field = cv2.GaussianBlur(heat_field, (0, 0), sigmaX=7, sigmaY=7)
        heat_norm = cv2.normalize(heat_field, None, 0, 255, cv2.NORM_MINMAX).astype(np.uint8)
        color_map = cv2.applyColorMap(heat_norm, cv2.COLORMAP_JET)
        alpha = np.clip(heat_norm.astype(np.float32) / 255.0 * 0.58, 0.0, 0.58)
        alpha = alpha[..., None]
        blended = image.astype(np.float32) * (1.0 - alpha) + color_map.astype(np.float32) * alpha
        return np.clip(blended, 0, 255).astype(np.uint8)

    def _apply_postprocess_to_payload(self, payload: Dict) -> None:
        if payload.get("_postprocessed"):
            return
        rows = payload.get("rows", []) or []
        annotated = payload.get("annotated")
        original = payload.get("original")
        alerts = []
        if annotated is not None:
            payload["_clean_annotated"] = annotated.copy()
            annotated = annotated.copy()
            payload["annotated"] = annotated

        rows = self._filter_rows_by_class_threshold(rows)
        if rows != payload.get("rows", []):
            payload["rows"] = rows
            payload["box_count"] = len(rows)
            if original is not None:
                annotated = self._draw_rows_on_image(original, rows)
                payload["_clean_annotated"] = annotated.copy()
                payload["annotated"] = annotated

        if self.heatmap_checkbox.isChecked() and annotated is not None:
            annotated = self._apply_detection_heatmap(annotated, rows)
            payload["annotated"] = annotated

        roi_hits = 0
        roi_rect = self._roi_rect_for_image(original) if self.roi_alert_checkbox.isChecked() else None
        if roi_rect and annotated is not None:
            x1, y1, x2, y2 = roi_rect
            cv2.rectangle(annotated, (x1, y1), (x2, y2), (0, 220, 255), 2)
            for row in rows:
                cx, cy = row.get("center", [None, None])
                if cx is not None and x1 <= float(cx) <= x2 and y1 <= float(cy) <= y2:
                    roi_hits += 1
            if roi_hits:
                alerts.append({"code": "ROI", "text": f"ROI报警：重点区域内检测到 {roi_hits} 个目标"})

        if self.alert_mark_checkbox.isChecked():
            low_conf_rows = [row for row in rows if float(row.get("confidence", 0.0) or 0.0) < self.low_conf_spin.value()]
            if low_conf_rows:
                alerts.append(
                    {
                        "code": "LOW CONF",
                        "text": f"低置信异常：{len(low_conf_rows)} 个目标置信度低于阈值 {self.low_conf_spin.value():.2f}",
                    }
                )
            current_count = int(payload.get("box_count", 0) or 0)
            if self.last_box_count_for_alert is not None:
                diff = abs(current_count - self.last_box_count_for_alert)
                if diff >= self.box_jump_spin.value():
                    alerts.append(
                        {
                            "code": "COUNT JUMP",
                            "text": f"数量突变异常：本次目标数 {current_count}，与上一次显示帧相差 {diff} 个",
                        }
                    )
            self.last_box_count_for_alert = current_count

        if self.alarm_rules_checkbox.isChecked():
            current_count = int(payload.get("box_count", 0) or 0)
            if current_count == 0:
                self.empty_frame_streak += 1
            else:
                self.empty_frame_streak = 0
            if self.empty_frame_streak >= self.no_target_spin.value():
                alerts.append(
                    {
                        "code": "NO TARGET",
                        "text": f"连续无目标报警：已连续 {self.empty_frame_streak} 个显示帧未检测到目标",
                    }
                )
            class_counter = Counter(row.get("class_name", "") for row in rows)
            for class_name, count in class_counter.items():
                if count >= self.class_count_spin.value():
                    alerts.append(
                        {
                            "code": "CLASS COUNT",
                            "text": f"类别数量报警：{class_name} 数量达到 {count} 个",
                        }
                    )

        if alerts and annotated is not None:
            cv2.rectangle(annotated, (0, 0), (min(annotated.shape[1] - 1, 520), 34), (0, 0, 180), -1)
            cv2.putText(
                annotated,
                "ALERT: " + " | ".join(alert["code"] for alert in alerts),
                (10, 24),
                cv2.FONT_HERSHEY_SIMPLEX,
                0.65,
                (255, 255, 255),
                2,
                cv2.LINE_AA,
            )

        payload["roi_hits"] = roi_hits
        payload["alerts"] = [alert["text"] for alert in alerts]
        payload["alert_codes"] = [alert["code"] for alert in alerts]
        payload["_postprocessed"] = True

    def _draw_roi_overlay(self, image, original, rows: Optional[List[Dict]] = None):
        if image is None or original is None or not self.roi_alert_checkbox.isChecked():
            return image
        output = image.copy()
        roi_rect = self._roi_rect_for_image(original)
        if not roi_rect:
            return output
        ox1, oy1, ox2, oy2 = roi_rect
        original_h, original_w = original.shape[:2]
        output_h, output_w = output.shape[:2]
        sx = output_w / max(original_w, 1)
        sy = output_h / max(original_h, 1)
        x1, y1, x2, y2 = int(ox1 * sx), int(oy1 * sy), int(ox2 * sx), int(oy2 * sy)
        cv2.rectangle(output, (x1, y1), (x2, y2), (0, 220, 255), 2)
        hits = 0
        for row in rows or []:
            cx, cy = row.get("center", [None, None])
            if cx is not None and ox1 <= float(cx) <= ox2 and oy1 <= float(cy) <= oy2:
                hits += 1
        label = f"ROI {self.roi_x_spin.value()}%,{self.roi_y_spin.value()}%,{self.roi_w_spin.value()}%,{self.roi_h_spin.value()}%"
        if hits:
            label += f"  Hits:{hits}"
        cv2.rectangle(output, (x1, max(0, y1 - 28)), (min(output_w - 1, x1 + 310), y1), (0, 160, 190), -1)
        cv2.putText(output, label, (x1 + 6, max(18, y1 - 8)), cv2.FONT_HERSHEY_SIMPLEX, 0.55, (255, 255, 255), 1, cv2.LINE_AA)
        return output

    def _refresh_roi_preview(self) -> None:
        if self.last_clean_annotated_image is None:
            return
        preview = self._draw_roi_overlay(self.last_clean_annotated_image, self.last_original_image, self.last_rows)
        self.last_annotated_image = preview
        self._refresh_image_panels(self.last_original_image, self.last_annotated_image)

    def _capture_alert_frame(self, payload: Dict) -> None:
        image = payload.get("annotated")
        if image is None:
            return
        source = str(payload.get("source_name", "source")).replace("\\", "_").replace("/", "_")
        stem = Path(source).stem or "source"
        frame_index = int(payload.get("frame_index", 0) or 0)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        image_path = os.path.join(self.alert_capture_dir, f"{timestamp}_{stem}_frame_{frame_index}.jpg")
        meta_path = os.path.splitext(image_path)[0] + ".json"
        cv2.imwrite(image_path, image)
        meta = {
            "captured_at": datetime.now().isoformat(timespec="seconds"),
            "source": payload.get("source_name", ""),
            "frame_index": frame_index,
            "alerts": payload.get("alerts", []),
            "rows": payload.get("rows", []),
        }
        with open(meta_path, "w", encoding="utf-8") as f:
            json.dump(meta, f, ensure_ascii=False, indent=2)

    # ----------------------- Worker callbacks -----------------------
    def _handle_frame_payload(self, payload: Dict) -> None:
        self._apply_postprocess_to_payload(payload)
        self.last_original_image = payload["original"]
        self.last_clean_annotated_image = payload.get("_clean_annotated", payload["annotated"]).copy()
        self.last_annotated_image = payload["annotated"]
        self.last_rows = payload["rows"]

        self._refresh_image_panels(payload["original"], payload["annotated"])
        self.box_count_card.setText(f"检测框数量: {payload['box_count']}")
        self._set_fps_value(payload["fps"])
        summary = self._build_summary(payload["rows"], payload["source_name"], payload["frame_index"])
        if payload.get("alerts"):
            summary_html = self._build_summary_html(summary, payload["alerts"])
            self.alert_records.append(
                {
                    "source": payload["source_name"],
                    "frame_index": payload["frame_index"],
                    "alerts": list(payload["alerts"]),
                }
            )
            if self.auto_capture_checkbox.isChecked():
                self._capture_alert_frame(payload)
        else:
            summary_html = self._build_summary_html(summary, [])
        self.summary_label.setText(summary_html)
        self._fill_table(payload["rows"])
        self._update_progress(payload["progress_current"], payload["progress_total"])
        self._append_session_records(payload)
        if self.backend_session_id is not None:
            self.backend.log_frame_async(self.backend_session_id, payload)
        self.state_source_value.setText(payload["source_name"])
        self.statusBar().showMessage(f"已处理: {payload['source_name']}")
        if getattr(self, "analysis_dialog", None) is not None and self.analysis_dialog.isVisible():
            self._refresh_current_analysis_tab()
            self._refresh_chart_tab()

    def _handle_worker_finished(self, message: str) -> None:
        self.backend.finish_session(self.backend_session_id, status="finished", message=message)
        self.is_running = False
        self.is_paused = False
        self.pause_button.setText("暂停")
        self.start_button.setEnabled(True)
        self.pause_button.setEnabled(False)
        self.export_button.setEnabled(bool(self.session_records))
        self.state_status_value.setText(message)
        self.statusBar().showMessage(message)

        if self.current_mode == DetectionMode.CAMERA:
            self.save_button.setEnabled(True)
            self.save_button.setText("开始录制")
            self.is_recording = False
        else:
            has_output = (
                (self.current_mode == DetectionMode.IMAGE and os.path.exists(self.temp_image_path))
                or (
                    self.current_mode == DetectionMode.FOLDER
                    and (self._folder_result_count() > 0 or self.last_annotated_image is not None)
                )
                or (self.current_mode == DetectionMode.VIDEO and os.path.exists(self.temp_video_path))
            )
            self.save_button.setEnabled(has_output)
            self.save_button.setText("保存结果")
        self._refresh_analysis_center()
        self._finish_current_queue_item("完成")

    def _handle_worker_error(self, message: str) -> None:
        self.backend.finish_session(self.backend_session_id, status="error", message=message)
        self.is_running = False
        self.is_paused = False
        self.start_button.setEnabled(True)
        self.pause_button.setEnabled(False)
        self.save_button.setEnabled(False)
        self.state_status_value.setText("异常")
        self.statusBar().showMessage("检测异常")
        self._refresh_analysis_center()
        self._finish_current_queue_item("失败")
        QMessageBox.critical(self, "错误", message)

    def _handle_recording_ready(self, temp_path: str) -> None:
        save_path, _ = QFileDialog.getSaveFileName(self, "保存摄像头录像", "", "视频文件 (*.mp4)")
        if not save_path:
            return
        try:
            safe_copy(temp_path, save_path)
            self.statusBar().showMessage(f"录像已保存到: {save_path}")
        except Exception as exc:
            QMessageBox.critical(self, "保存失败", str(exc))

    def _set_status(self, text: str) -> None:
        self.state_status_value.setText(text)
        self.statusBar().showMessage(text)

    def _folder_result_count(self) -> int:
        if not os.path.isdir(self.temp_folder_dir):
            return 0
        return sum(1 for name in os.listdir(self.temp_folder_dir) if os.path.isfile(os.path.join(self.temp_folder_dir, name)))

    def _start_backend_session(self) -> None:
        source = ""
        if self.current_mode == DetectionMode.IMAGE:
            source = self.current_image_path or ""
        elif self.current_mode == DetectionMode.FOLDER:
            source = self.folder_path_label.toolTip() or self.folder_path_label.text()
        elif self.current_mode == DetectionMode.VIDEO:
            source = self.current_video_path or ""
        elif self.current_mode == DetectionMode.CAMERA:
            source = f"camera:{self.current_camera_id}"

        model_path = self.detector.model_path or ""
        self.backend_session_id = self.backend.start_session(
            mode=self.current_mode.name.lower(),
            source=source,
            model_name=Path(model_path).name,
            model_path=model_path,
            runtime=self.detector.runtime_label(),
            imgsz=self._current_imgsz(),
            conf=self.conf_spin.value(),
            iou=self.iou_spin.value(),
        )
        self.statusBar().showMessage(f"本地后端会话已创建：#{self.backend_session_id}")

    # ----------------------- Analysis center -----------------------
    def _update_resource_monitor(self) -> None:
        if psutil is not None:
            cpu = psutil.cpu_percent(interval=None)
            memory = psutil.virtual_memory()
            process_memory = psutil.Process(os.getpid()).memory_info().rss / 1024 / 1024
            cpu_text = f"CPU {cpu:.0f}%"
            memory_text = f"{memory.percent:.0f}% / {process_memory:.0f}MB"
        else:
            cpu_text = "CPU --"
            memory_text = "--"
        if self.is_running:
            gpu_text = self.state_gpu_value.text() or "GPU --"
            vram_text = self.state_vram_value.text() or "--"
        else:
            gpu_text, vram_text = self._query_gpu_usage()
        self.state_resource_value.setText(cpu_text)
        self.state_memory_value.setText(memory_text)
        self.state_gpu_value.setText(gpu_text)
        self.state_vram_value.setText(vram_text)

        if getattr(self, "analysis_dialog", None) is not None and self.analysis_dialog.isVisible():
            self._refresh_current_analysis_tab()

    def _query_gpu_usage(self):
        try:
            output = subprocess.check_output(
                [
                    "nvidia-smi",
                    "--query-gpu=utilization.gpu,memory.used,memory.total",
                    "--format=csv,noheader,nounits",
                ],
                text=True,
                timeout=0.8,
                stderr=subprocess.DEVNULL,
            )
            first_line = output.strip().splitlines()[0]
            util, used, total = [part.strip() for part in first_line.split(",")[:3]]
            return f"GPU {util}%", f"{used} / {total}MB"
        except Exception:
            return "GPU 不可用", "--"

    def open_analysis_center(self) -> None:
        if getattr(self, "analysis_dialog", None) is None:
            self._build_analysis_center()
        self._refresh_analysis_center()
        self.analysis_dialog.show()
        self.analysis_dialog.raise_()
        self.analysis_dialog.activateWindow()

    def _build_analysis_center(self) -> None:
        self.analysis_dialog = QDialog(self)
        self.analysis_dialog.setWindowTitle("系统工作台 / 实验分析中心")
        self.analysis_dialog.resize(980, 680)

        tabs = QTabWidget()

        history_tab = QWidget()
        history_layout = QVBoxLayout(history_tab)
        history_buttons = QHBoxLayout()
        refresh_button = QPushButton("刷新历史")
        refresh_button.clicked.connect(self._refresh_analysis_center)
        report_button = QPushButton("生成报告")
        report_button.clicked.connect(self.export_analysis_report)
        word_report_button = QPushButton("生成Word报告")
        word_report_button.clicked.connect(self.export_word_report)
        history_buttons.addWidget(refresh_button)
        history_buttons.addWidget(report_button)
        history_buttons.addWidget(word_report_button)
        history_buttons.addStretch(1)

        self.analysis_history_table = QTableWidget(0, 10)
        self.analysis_history_table.setHorizontalHeaderLabels(
            ["ID", "开始时间", "模式", "模型", "后端", "尺寸", "状态", "采样帧", "目标数", "平均FPS"]
        )
        self.analysis_history_table.verticalHeader().setVisible(False)
        self.analysis_history_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.analysis_history_table.setSelectionBehavior(QTableWidget.SelectRows)
        self.analysis_history_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        self.analysis_history_table.horizontalHeader().setSectionResizeMode(4, QHeaderView.Stretch)
        history_layout.addLayout(history_buttons)
        history_layout.addWidget(self.analysis_history_table)

        current_tab = QWidget()
        current_layout = QVBoxLayout(current_tab)
        self.analysis_current_text = QTextEdit()
        self.analysis_current_text.setReadOnly(True)
        current_layout.addWidget(self.analysis_current_text)

        charts_tab = QWidget()
        charts_layout = QVBoxLayout(charts_tab)
        chart_buttons = QHBoxLayout()
        refresh_charts_button = QPushButton("刷新图表")
        refresh_charts_button.clicked.connect(self._refresh_chart_tab)
        export_chart_button = QPushButton("导出图表PNG")
        export_chart_button.clicked.connect(self.export_statistics_charts)
        chart_buttons.addWidget(refresh_charts_button)
        chart_buttons.addWidget(export_chart_button)
        chart_buttons.addStretch(1)
        self.class_chart_label = QLabel("暂无类别统计图")
        self.class_chart_label.setAlignment(Qt.AlignCenter)
        self.timeline_chart_label = QLabel("暂无帧级趋势图")
        self.timeline_chart_label.setAlignment(Qt.AlignCenter)
        charts_layout.addLayout(chart_buttons)
        charts_layout.addWidget(self.class_chart_label)
        charts_layout.addWidget(self.timeline_chart_label)

        benchmark_tab = QWidget()
        benchmark_layout = QVBoxLayout(benchmark_tab)
        benchmark_controls = QHBoxLayout()
        sample_label = QLabel("样本上限")
        self.benchmark_limit_spin = QSpinBox()
        self.benchmark_limit_spin.setRange(1, 1000)
        self.benchmark_limit_spin.setValue(self.config.analysis_benchmark_limit)
        run_benchmark_button = QPushButton("运行尺寸对比")
        run_benchmark_button.clicked.connect(self.run_parameter_benchmark)
        export_benchmark_button = QPushButton("导出对比CSV")
        export_benchmark_button.clicked.connect(self.export_benchmark_csv)
        self.benchmark_run_button = run_benchmark_button
        self.benchmark_export_button = export_benchmark_button
        benchmark_controls.addWidget(sample_label)
        benchmark_controls.addWidget(self.benchmark_limit_spin)
        benchmark_controls.addWidget(run_benchmark_button)
        benchmark_controls.addWidget(export_benchmark_button)
        benchmark_controls.addStretch(1)

        self.benchmark_status_label = QLabel("选择图片/文件夹/视频后，可自动测试 320、416、640 输入尺寸。")
        self.benchmark_table = QTableWidget(0, 7)
        self.benchmark_table.setHorizontalHeaderLabels(
            ["输入尺寸", "样本数", "总耗时(s)", "FPS", "平均耗时(ms)", "目标数", "平均置信度"]
        )
        self.benchmark_table.verticalHeader().setVisible(False)
        self.benchmark_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.benchmark_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        benchmark_layout.addLayout(benchmark_controls)
        benchmark_layout.addWidget(self.benchmark_status_label)
        benchmark_layout.addWidget(self.benchmark_table)

        thresholds_tab = QWidget()
        thresholds_layout = QVBoxLayout(thresholds_tab)
        threshold_buttons = QHBoxLayout()
        add_threshold_button = QPushButton("添加/更新类别阈值")
        add_threshold_button.clicked.connect(self.add_or_update_class_threshold)
        delete_threshold_button = QPushButton("删除选中阈值")
        delete_threshold_button.clicked.connect(self.delete_selected_class_threshold)
        clear_threshold_button = QPushButton("清空阈值")
        clear_threshold_button.clicked.connect(self.clear_class_thresholds)
        threshold_buttons.addWidget(add_threshold_button)
        threshold_buttons.addWidget(delete_threshold_button)
        threshold_buttons.addWidget(clear_threshold_button)
        threshold_buttons.addStretch(1)
        self.class_threshold_table = QTableWidget(0, 2)
        self.class_threshold_table.setHorizontalHeaderLabels(["类别", "最低置信度"])
        self.class_threshold_table.verticalHeader().setVisible(False)
        self.class_threshold_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.class_threshold_table.setSelectionBehavior(QTableWidget.SelectRows)
        self.class_threshold_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        threshold_note = QLabel("说明：类别阈值会过滤表格、摘要、日志和导出结果；若启用后检测图会按过滤结果重新绘制。")
        threshold_note.setWordWrap(True)
        thresholds_layout.addLayout(threshold_buttons)
        thresholds_layout.addWidget(threshold_note)
        thresholds_layout.addWidget(self.class_threshold_table)

        queue_tab = QWidget()
        queue_layout = QVBoxLayout(queue_tab)
        queue_buttons = QHBoxLayout()
        add_queue_image_button = QPushButton("加入图片")
        add_queue_image_button.clicked.connect(self.add_queue_image)
        add_queue_folder_button = QPushButton("加入文件夹")
        add_queue_folder_button.clicked.connect(self.add_queue_folder)
        add_queue_video_button = QPushButton("加入视频")
        add_queue_video_button.clicked.connect(self.add_queue_video)
        add_queue_stream_button = QPushButton("加入远程流")
        add_queue_stream_button.clicked.connect(self.add_queue_remote_stream)
        run_queue_button = QPushButton("运行队列")
        run_queue_button.clicked.connect(self.run_task_queue)
        clear_queue_button = QPushButton("清空队列")
        clear_queue_button.clicked.connect(self.clear_task_queue)
        self.run_queue_button = run_queue_button
        queue_buttons.addWidget(add_queue_image_button)
        queue_buttons.addWidget(add_queue_folder_button)
        queue_buttons.addWidget(add_queue_video_button)
        queue_buttons.addWidget(add_queue_stream_button)
        queue_buttons.addWidget(run_queue_button)
        queue_buttons.addWidget(clear_queue_button)
        queue_buttons.addStretch(1)
        self.queue_status_label = QLabel("任务队列可批量执行图片、文件夹、视频和远程流检测。")
        self.queue_table = QTableWidget(0, 4)
        self.queue_table.setHorizontalHeaderLabels(["序号", "类型", "来源", "状态"])
        self.queue_table.verticalHeader().setVisible(False)
        self.queue_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.queue_table.horizontalHeader().setSectionResizeMode(2, QHeaderView.Stretch)
        queue_layout.addLayout(queue_buttons)
        queue_layout.addWidget(self.queue_status_label)
        queue_layout.addWidget(self.queue_table)

        search_tab = QWidget()
        search_layout = QVBoxLayout(search_tab)
        search_controls = QHBoxLayout()
        self.search_keyword_edit = QLineEdit()
        self.search_keyword_edit.setPlaceholderText("输入类别/来源/模型/状态关键词")
        self.search_limit_spin = QSpinBox()
        self.search_limit_spin.setRange(10, 5000)
        self.search_limit_spin.setValue(300)
        search_button = QPushButton("检索")
        search_button.clicked.connect(self.search_history_results)
        export_search_button = QPushButton("导出检索CSV")
        export_search_button.clicked.connect(self.export_search_results_csv)
        self.search_export_button = export_search_button
        self.search_export_button.setEnabled(False)
        search_controls.addWidget(self.search_keyword_edit)
        search_controls.addWidget(QLabel("上限"))
        search_controls.addWidget(self.search_limit_spin)
        search_controls.addWidget(search_button)
        search_controls.addWidget(export_search_button)
        self.search_table = QTableWidget(0, 9)
        self.search_table.setHorizontalHeaderLabels(["会话", "时间", "模式", "来源", "模型", "帧", "类别", "置信度", "状态"])
        self.search_table.verticalHeader().setVisible(False)
        self.search_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.search_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        self.search_table.horizontalHeader().setSectionResizeMode(3, QHeaderView.Stretch)
        self.search_rows: List[Dict] = []
        search_layout.addLayout(search_controls)
        search_layout.addWidget(self.search_table)

        presets_tab = QWidget()
        presets_layout = QVBoxLayout(presets_tab)
        preset_buttons = QHBoxLayout()
        save_preset_button = QPushButton("保存当前为预设")
        save_preset_button.clicked.connect(self.save_current_preset)
        apply_preset_button = QPushButton("应用选中预设")
        apply_preset_button.clicked.connect(self.apply_selected_preset)
        delete_preset_button = QPushButton("删除选中预设")
        delete_preset_button.clicked.connect(self.delete_selected_preset)
        refresh_preset_button = QPushButton("刷新预设")
        refresh_preset_button.clicked.connect(self._populate_preset_table)
        preset_buttons.addWidget(save_preset_button)
        preset_buttons.addWidget(apply_preset_button)
        preset_buttons.addWidget(delete_preset_button)
        preset_buttons.addWidget(refresh_preset_button)
        preset_buttons.addStretch(1)
        self.preset_table = QTableWidget(0, 8)
        self.preset_table.setHorizontalHeaderLabels(["名称", "置信度", "IoU", "尺寸", "原图对比", "ROI", "低置信", "数量突变"])
        self.preset_table.verticalHeader().setVisible(False)
        self.preset_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.preset_table.setSelectionBehavior(QTableWidget.SelectRows)
        self.preset_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        self.preset_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        presets_layout.addLayout(preset_buttons)
        presets_layout.addWidget(self.preset_table)

        tabs.addTab(history_tab, "实验历史")
        tabs.addTab(current_tab, "当前统计")
        tabs.addTab(charts_tab, "统计图表")
        tabs.addTab(benchmark_tab, "参数对比")
        tabs.addTab(thresholds_tab, "类别阈值")
        tabs.addTab(queue_tab, "任务队列")
        tabs.addTab(search_tab, "历史检索")
        tabs.addTab(presets_tab, "参数预设")

        root_layout = QVBoxLayout(self.analysis_dialog)
        root_layout.addWidget(tabs)

    def _refresh_analysis_center(self) -> None:
        if getattr(self, "analysis_dialog", None) is None:
            return
        self._refresh_history_table()
        self._refresh_current_analysis_tab()
        self._refresh_chart_tab()
        self._populate_benchmark_table()
        self._populate_class_threshold_table()
        self._populate_queue_table()
        self._populate_preset_table()

    def _refresh_history_table(self) -> None:
        sessions = self.backend.list_sessions()
        self.analysis_history_table.setRowCount(len(sessions))
        for row_index, session in enumerate(sessions):
            values = [
                session.get("id", ""),
                session.get("started_at", ""),
                session.get("mode", ""),
                session.get("model_name", ""),
                session.get("runtime", ""),
                session.get("imgsz", ""),
                session.get("status", ""),
                session.get("frame_count", 0),
                session.get("detection_count", 0),
                self._format_float(session.get("avg_fps"), 1),
            ]
            for col, value in enumerate(values):
                self.analysis_history_table.setItem(row_index, col, QTableWidgetItem(str(value)))

    def _refresh_current_analysis_tab(self) -> None:
        if getattr(self, "analysis_current_text", None) is None:
            return
        self.analysis_current_text.setPlainText(self._build_current_stats_text())

    def _refresh_chart_tab(self) -> None:
        if getattr(self, "class_chart_label", None) is None:
            return
        class_chart = self._build_class_distribution_chart()
        timeline_chart = self._build_timeline_chart()
        self.class_chart_label.setPixmap(class_chart)
        self.timeline_chart_label.setPixmap(timeline_chart)

    def _build_class_distribution_chart(self, width: int = 880, height: int = 250) -> QPixmap:
        pixmap = QPixmap(width, height)
        pixmap.fill(QColor("#ffffff"))
        painter = QPainter(pixmap)
        try:
            painter.setRenderHint(QPainter.Antialiasing)
            painter.setPen(QPen(QColor("#17324a")))
            painter.setFont(QFont("SimSun", 12, QFont.Bold))
            painter.drawText(18, 28, "类别数量分布")
            counter = Counter(record["class_name"] for record in self.session_records)
            if not counter:
                painter.setFont(QFont("SimSun", 11))
                painter.drawText(18, 90, "暂无检测类别数据。运行图片、文件夹、视频或摄像头检测后会自动生成统计图。")
                return pixmap
            items = counter.most_common(8)
            max_count = max(count for _, count in items) or 1
            bar_left = 130
            bar_top = 55
            bar_height = 18
            gap = 10
            painter.setFont(QFont("SimSun", 10))
            for idx, (name, count) in enumerate(items):
                y = bar_top + idx * (bar_height + gap)
                painter.setPen(QPen(QColor("#17324a")))
                painter.drawText(18, y + 14, str(name)[:14])
                bar_width = int((width - bar_left - 110) * count / max_count)
                painter.fillRect(bar_left, y, bar_width, bar_height, QColor("#2f6f9f"))
                painter.setPen(QPen(QColor("#1f4f7a")))
                painter.drawRect(bar_left, y, max(1, bar_width), bar_height)
                painter.drawText(bar_left + bar_width + 8, y + 14, str(count))
        finally:
            painter.end()
        return pixmap

    def _build_timeline_chart(self, width: int = 880, height: int = 250) -> QPixmap:
        pixmap = QPixmap(width, height)
        pixmap.fill(QColor("#ffffff"))
        painter = QPainter(pixmap)
        try:
            painter.setRenderHint(QPainter.Antialiasing)
            painter.setPen(QPen(QColor("#17324a")))
            painter.setFont(QFont("SimSun", 12, QFont.Bold))
            painter.drawText(18, 28, "帧级趋势（FPS 与目标数量）")
            samples = self.frame_samples[-80:]
            if not samples:
                painter.setFont(QFont("SimSun", 11))
                painter.drawText(18, 90, "暂无帧级采样数据。")
                return pixmap
            left, top, right, bottom = 55, 55, width - 35, height - 35
            painter.setPen(QPen(QColor("#7aa2c5"), 1))
            painter.drawRect(left, top, right - left, bottom - top)
            fps_values = [float(sample.get("fps", 0.0) or 0.0) for sample in samples]
            box_values = [int(sample.get("box_count", 0) or 0) for sample in samples]
            max_fps = max(fps_values) or 1.0
            max_boxes = max(box_values) or 1
            n = max(1, len(samples) - 1)
            points = []
            for idx, fps in enumerate(fps_values):
                x = left + int((right - left) * idx / n)
                y = bottom - int((bottom - top) * fps / max_fps)
                points.append((x, y))
            painter.setPen(QPen(QColor("#2f6f9f"), 2))
            for p1, p2 in zip(points, points[1:]):
                painter.drawLine(p1[0], p1[1], p2[0], p2[1])
            bar_w = max(2, int((right - left) / max(1, len(samples)) * 0.55))
            painter.setPen(QPen(QColor("#c37a1c")))
            painter.setBrush(QColor("#f2b86d"))
            for idx, count in enumerate(box_values):
                x = left + int((right - left) * idx / max(1, len(samples)))
                h = int((bottom - top) * count / max_boxes)
                painter.drawRect(x, bottom - h, bar_w, h)
            painter.setPen(QPen(QColor("#17324a")))
            painter.setFont(QFont("SimSun", 9))
            painter.drawText(left, top - 8, f"FPS最高 {max_fps:.1f}")
            painter.drawText(right - 120, top - 8, f"目标最高 {max_boxes}")
            painter.drawText(left, bottom + 20, "蓝线: FPS")
            painter.drawText(left + 95, bottom + 20, "橙柱: 目标数量")
        finally:
            painter.end()
        return pixmap

    def export_statistics_charts(self) -> None:
        save_dir = QFileDialog.getExistingDirectory(self, "选择图表保存目录")
        if not save_dir:
            return
        class_path = os.path.join(save_dir, "class_distribution.png")
        timeline_path = os.path.join(save_dir, "frame_timeline.png")
        self._build_class_distribution_chart().save(class_path)
        self._build_timeline_chart().save(timeline_path)
        self.statusBar().showMessage(f"统计图表已导出到: {save_dir}")

    def _preset_store(self) -> Dict:
        raw = self.settings.value("presets/catalog", "{}")
        try:
            data = json.loads(raw)
            return data if isinstance(data, dict) else {}
        except Exception:
            return {}

    def _save_preset_store(self, presets: Dict) -> None:
        self.settings.setValue("presets/catalog", json.dumps(presets, ensure_ascii=False))

    def _current_preset_data(self) -> Dict:
        return {
            "conf": self.conf_spin.value(),
            "iou": self.iou_spin.value(),
            "imgsz": self._current_imgsz(),
            "show_labels": self.show_labels_checkbox.isChecked(),
            "show_conf": self.show_conf_checkbox.isChecked(),
            "compare_view": self.compare_view_checkbox.isChecked(),
            "heatmap": self.heatmap_checkbox.isChecked(),
            "save_folder_results": self.save_folder_results_checkbox.isChecked(),
            "roi_alert": self.roi_alert_checkbox.isChecked(),
            "alert_mark": self.alert_mark_checkbox.isChecked(),
            "auto_capture": self.auto_capture_checkbox.isChecked(),
            "alarm_rules": self.alarm_rules_checkbox.isChecked(),
            "low_conf": self.low_conf_spin.value(),
            "box_jump": self.box_jump_spin.value(),
            "no_target_frames": self.no_target_spin.value(),
            "class_count_alert": self.class_count_spin.value(),
            "roi": [
                self.roi_x_spin.value(),
                self.roi_y_spin.value(),
                self.roi_w_spin.value(),
                self.roi_h_spin.value(),
            ],
        }

    def save_current_preset(self) -> None:
        name, ok = QInputDialog.getText(self, "保存参数预设", "请输入预设名称：")
        if not ok or not name.strip():
            return
        presets = self._preset_store()
        presets[name.strip()] = self._current_preset_data()
        self._save_preset_store(presets)
        self._populate_preset_table()
        self.statusBar().showMessage(f"参数预设已保存：{name.strip()}")

    def _selected_preset_name(self) -> Optional[str]:
        if getattr(self, "preset_table", None) is None:
            return None
        selected = self.preset_table.selectionModel().selectedRows()
        if not selected:
            return None
        return self.preset_table.item(selected[0].row(), 0).text()

    def apply_selected_preset(self) -> None:
        name = self._selected_preset_name()
        if not name:
            QMessageBox.information(self, "提示", "请先选择一个参数预设。")
            return
        preset = self._preset_store().get(name)
        if not preset:
            return
        self.conf_spin.setValue(float(preset.get("conf", self.config.default_confidence)))
        self.iou_spin.setValue(float(preset.get("iou", self.config.default_iou)))
        self.imgsz_spin.setValue(int(preset.get("imgsz", self.config.inference_imgsz)))
        self.show_labels_checkbox.setChecked(bool(preset.get("show_labels", True)))
        self.show_conf_checkbox.setChecked(bool(preset.get("show_conf", True)))
        self.compare_view_checkbox.setChecked(bool(preset.get("compare_view", False)))
        self.heatmap_checkbox.setChecked(bool(preset.get("heatmap", False)))
        self.save_folder_results_checkbox.setChecked(bool(preset.get("save_folder_results", False)))
        self.roi_alert_checkbox.setChecked(bool(preset.get("roi_alert", False)))
        self.alert_mark_checkbox.setChecked(bool(preset.get("alert_mark", True)))
        self.auto_capture_checkbox.setChecked(bool(preset.get("auto_capture", True)))
        self.alarm_rules_checkbox.setChecked(bool(preset.get("alarm_rules", False)))
        self.low_conf_spin.setValue(float(preset.get("low_conf", self.config.default_low_confidence_alert)))
        self.box_jump_spin.setValue(int(preset.get("box_jump", self.config.default_box_jump_alert)))
        self.no_target_spin.setValue(int(preset.get("no_target_frames", self.config.default_no_target_alert_frames)))
        self.class_count_spin.setValue(int(preset.get("class_count_alert", self.config.default_class_count_alert)))
        roi = preset.get("roi", self.config.default_roi_percent)
        for spin, value in zip([self.roi_x_spin, self.roi_y_spin, self.roi_w_spin, self.roi_h_spin], roi):
            spin.setValue(int(value))
        self._save_settings()
        self.statusBar().showMessage(f"已应用参数预设：{name}")

    def delete_selected_preset(self) -> None:
        name = self._selected_preset_name()
        if not name:
            QMessageBox.information(self, "提示", "请先选择一个参数预设。")
            return
        presets = self._preset_store()
        presets.pop(name, None)
        self._save_preset_store(presets)
        self._populate_preset_table()
        self.statusBar().showMessage(f"已删除参数预设：{name}")

    def _populate_preset_table(self) -> None:
        if getattr(self, "preset_table", None) is None:
            return
        presets = self._preset_store()
        self.preset_table.setRowCount(len(presets))
        for row, (name, preset) in enumerate(sorted(presets.items())):
            values = [
                name,
                f"{float(preset.get('conf', 0.0)):.2f}",
                f"{float(preset.get('iou', 0.0)):.2f}",
                preset.get("imgsz", ""),
                "是" if preset.get("compare_view", False) else "否",
                "是" if preset.get("roi_alert", False) else "否",
                f"{float(preset.get('low_conf', self.config.default_low_confidence_alert)):.2f}",
                preset.get("box_jump", self.config.default_box_jump_alert),
            ]
            for col, value in enumerate(values):
                self.preset_table.setItem(row, col, QTableWidgetItem(str(value)))

    def _load_class_thresholds(self) -> Dict[str, float]:
        raw = self.settings.value("class_thresholds/catalog", "{}")
        try:
            data = json.loads(raw)
            if isinstance(data, dict):
                return {str(k): float(v) for k, v in data.items()}
        except Exception:
            pass
        return {}

    def _save_class_thresholds(self) -> None:
        self.settings.setValue("class_thresholds/catalog", json.dumps(self.class_thresholds, ensure_ascii=False))

    def add_or_update_class_threshold(self) -> None:
        if self.detector.is_loaded and self.detector.names:
            class_items = sorted(str(name) for name in self.detector.names.values())
            class_name, ok = QInputDialog.getItem(self, "类别阈值", "选择类别：", class_items, 0, True)
        else:
            class_name, ok = QInputDialog.getText(self, "类别阈值", "输入类别名称：")
        if not ok or not class_name.strip():
            return
        current = float(self.class_thresholds.get(class_name.strip(), self.conf_spin.value()))
        value, ok = QInputDialog.getDouble(self, "类别阈值", "最低置信度：", current, 0.01, 1.0, 2)
        if not ok:
            return
        self.class_thresholds[class_name.strip()] = float(value)
        self._save_class_thresholds()
        self._populate_class_threshold_table()
        self.statusBar().showMessage(f"已设置类别阈值：{class_name.strip()} >= {value:.2f}")

    def delete_selected_class_threshold(self) -> None:
        if getattr(self, "class_threshold_table", None) is None:
            return
        selected = self.class_threshold_table.selectionModel().selectedRows()
        if not selected:
            QMessageBox.information(self, "提示", "请先选择一个类别阈值。")
            return
        name = self.class_threshold_table.item(selected[0].row(), 0).text()
        self.class_thresholds.pop(name, None)
        self._save_class_thresholds()
        self._populate_class_threshold_table()

    def clear_class_thresholds(self) -> None:
        self.class_thresholds = {}
        self._save_class_thresholds()
        self._populate_class_threshold_table()
        self.statusBar().showMessage("类别阈值已清空")

    def _populate_class_threshold_table(self) -> None:
        if getattr(self, "class_threshold_table", None) is None:
            return
        items = sorted(self.class_thresholds.items())
        self.class_threshold_table.setRowCount(len(items))
        for row, (name, threshold) in enumerate(items):
            self.class_threshold_table.setItem(row, 0, QTableWidgetItem(str(name)))
            self.class_threshold_table.setItem(row, 1, QTableWidgetItem(f"{float(threshold):.2f}"))

    def _add_queue_item(self, mode: DetectionMode, source, label: str) -> None:
        self.task_queue.append({"mode": mode, "source": source, "label": label, "status": "等待"})
        self._populate_queue_table()

    def add_queue_image(self) -> None:
        paths, _ = QFileDialog.getOpenFileNames(self, "加入图片任务", "", "图片文件 (*.jpg *.jpeg *.png *.bmp *.webp *.tif *.tiff)")
        for path in paths:
            self._add_queue_item(DetectionMode.IMAGE, path, Path(path).name)

    def add_queue_folder(self) -> None:
        folder = QFileDialog.getExistingDirectory(self, "加入文件夹任务")
        if folder:
            self._add_queue_item(DetectionMode.FOLDER, folder, Path(folder).name)

    def add_queue_video(self) -> None:
        path, _ = QFileDialog.getOpenFileName(self, "加入视频任务", "", "视频文件 (*.mp4 *.avi *.mov *.mkv *.flv *.wmv)")
        if path:
            self._add_queue_item(DetectionMode.VIDEO, path, Path(path).name)

    def add_queue_remote_stream(self) -> None:
        url, ok = QInputDialog.getText(self, "加入远程流任务", "请输入 RTSP/HTTP/HTTPS 地址：")
        if ok and url.strip():
            self._add_queue_item(DetectionMode.VIDEO, url.strip(), url.strip())

    def clear_task_queue(self) -> None:
        if self.queue_running:
            QMessageBox.information(self, "提示", "队列运行中，无法清空。")
            return
        self.task_queue = []
        self.queue_index = 0
        self._populate_queue_table()

    def run_task_queue(self) -> None:
        if not self.detector.is_loaded:
            QMessageBox.warning(self, "提示", "请先加载模型。")
            return
        if not self.task_queue:
            QMessageBox.information(self, "提示", "请先加入检测任务。")
            return
        if self.is_running:
            QMessageBox.warning(self, "提示", "请等待当前检测任务结束。")
            return
        self.queue_running = True
        self.queue_index = 0
        self._run_next_queue_item()

    def _run_next_queue_item(self) -> None:
        if not self.queue_running:
            return
        if self.queue_index >= len(self.task_queue):
            self.queue_running = False
            self.queue_status_label.setText("任务队列已完成。")
            self._populate_queue_table()
            return
        item = self.task_queue[self.queue_index]
        item["status"] = "运行中"
        self._populate_queue_table()
        self.queue_status_label.setText(f"正在运行队列任务 {self.queue_index + 1}/{len(self.task_queue)}：{item['label']}")
        try:
            if item["mode"] == DetectionMode.IMAGE:
                self._set_image_source(str(item["source"]))
            elif item["mode"] == DetectionMode.FOLDER:
                self._set_folder_source(str(item["source"]))
            elif item["mode"] == DetectionMode.VIDEO:
                source = str(item["source"])
                self._set_video_source(source, remote=source.lower().startswith(("rtsp://", "http://", "https://")))
            self.start_detection()
        except Exception as exc:
            item["status"] = f"失败: {exc}"
            self.queue_index += 1
            QTimer.singleShot(100, self._run_next_queue_item)

    def _finish_current_queue_item(self, status: str) -> None:
        if not self.queue_running or self.queue_index >= len(self.task_queue):
            return
        self.task_queue[self.queue_index]["status"] = status
        self.queue_index += 1
        self._populate_queue_table()
        QTimer.singleShot(250, self._run_next_queue_item)

    def _populate_queue_table(self) -> None:
        if getattr(self, "queue_table", None) is None:
            return
        self.queue_table.setRowCount(len(self.task_queue))
        for row, item in enumerate(self.task_queue):
            values = [row + 1, item["mode"].name.lower(), item["label"], item.get("status", "等待")]
            for col, value in enumerate(values):
                self.queue_table.setItem(row, col, QTableWidgetItem(str(value)))
        if getattr(self, "run_queue_button", None) is not None:
            self.run_queue_button.setEnabled(bool(self.task_queue) and not self.queue_running)

    def search_history_results(self) -> None:
        keyword = self.search_keyword_edit.text().strip()
        self.search_rows = self.backend.search_results(keyword, self.search_limit_spin.value())
        self.search_table.setRowCount(len(self.search_rows))
        for row, item in enumerate(self.search_rows):
            values = [
                item.get("session_id", ""),
                item.get("started_at", ""),
                item.get("mode", ""),
                item.get("source_name") or item.get("source", ""),
                item.get("model_name", ""),
                item.get("frame_index", ""),
                item.get("class_name", ""),
                self._format_float(item.get("confidence"), 3),
                item.get("status", ""),
            ]
            for col, value in enumerate(values):
                self.search_table.setItem(row, col, QTableWidgetItem(str(value)))
        self.search_export_button.setEnabled(bool(self.search_rows))
        self.statusBar().showMessage(f"检索完成：{len(self.search_rows)} 条记录")

    def export_search_results_csv(self) -> None:
        if not getattr(self, "search_rows", []):
            QMessageBox.information(self, "提示", "当前没有可导出的检索结果。")
            return
        save_path, _ = QFileDialog.getSaveFileName(self, "导出检索结果", "search_results.csv", "CSV 文件 (*.csv)")
        if not save_path:
            return
        with open(save_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=list(self.search_rows[0].keys()))
            writer.writeheader()
            writer.writerows(self.search_rows)
        self.statusBar().showMessage(f"检索结果已导出到: {save_path}")

    def _build_current_stats_text(self) -> str:
        fps_values = [sample["fps"] for sample in self.frame_samples if sample.get("fps", 0) > 0]
        target_count = sum(int(sample.get("box_count", 0) or 0) for sample in self.frame_samples)
        class_counter = Counter(record["class_name"] for record in self.session_records)
        confidences = [float(record["confidence"]) for record in self.session_records]

        source = self._current_source_text()
        model_name = Path(self.detector.model_path).name if self.detector.model_path else "未加载"
        lines = [
            "当前实验概览",
            f"实验编号: {self.backend_session_id or '尚未创建'}",
            f"模型: {model_name}",
            f"运行后端: {self.detector.runtime_label() if self.detector.is_loaded else '未加载'}",
            f"检测模式: {self.current_mode.name.lower()}",
            f"数据来源: {source}",
            f"输入尺寸: {self._current_imgsz()}",
            f"置信度 / IoU: {self.conf_spin.value():.2f} / {self.iou_spin.value():.2f}",
            "",
            "性能与结果",
            f"采样帧数: {len(self.frame_samples)}",
            f"目标总数: {target_count}",
            f"平均FPS: {statistics.mean(fps_values):.2f}" if fps_values else "平均FPS: --",
            f"最高FPS: {max(fps_values):.2f}" if fps_values else "最高FPS: --",
            f"平均置信度: {statistics.mean(confidences):.3f}" if confidences else "平均置信度: --",
            f"最高置信度: {max(confidences):.3f}" if confidences else "最高置信度: --",
            f"异常帧数: {len(self.alert_records)}",
            f"ROI报警: {'开启' if self.roi_alert_checkbox.isChecked() else '关闭'}",
            "",
            "实时资源",
            f"CPU: {self.state_resource_value.text()}",
            f"内存: {self.state_memory_value.text()}",
            f"GPU: {self.state_gpu_value.text()}",
            f"显存: {self.state_vram_value.text()}",
            "",
            "类别分布",
        ]
        if class_counter:
            lines.extend(f"{name}: {count}" for name, count in class_counter.most_common())
        else:
            lines.append("暂无类别统计")
        if self.alert_records:
            lines.extend(["", "最近异常帧"])
            for alert in self.alert_records[-8:]:
                lines.append(f"{alert['source']}#{alert['frame_index']}: {'；'.join(alert['alerts'])}")
        return "\n".join(lines)

    def run_parameter_benchmark(self) -> None:
        if not self.detector.is_loaded:
            QMessageBox.warning(self, "提示", "请先加载模型。")
            return
        if self.is_running:
            QMessageBox.warning(self, "提示", "请先等待当前检测任务结束，再运行参数对比。")
            return
        if self.current_mode not in {DetectionMode.IMAGE, DetectionMode.FOLDER, DetectionMode.VIDEO}:
            QMessageBox.warning(self, "提示", "参数对比支持图片、文件夹和视频输入。")
            return

        source = self._benchmark_source()
        if source is None:
            QMessageBox.warning(self, "提示", "请先选择可用于测试的图片、文件夹或视频。")
            return

        self.benchmark_rows = []
        self._populate_benchmark_table()
        self.benchmark_status_label.setText("正在运行参数对比，请稍候...")
        self.benchmark_run_button.setEnabled(False)
        self.benchmark_export_button.setEnabled(False)

        self.benchmark_worker = BenchmarkWorker(
            detector=self.detector,
            mode=self.current_mode,
            source=source,
            sizes=self.config.analysis_benchmark_sizes,
            sample_limit=self.benchmark_limit_spin.value(),
            conf=self.conf_spin.value(),
            iou=self.iou_spin.value(),
        )
        self.benchmark_worker.status_changed.connect(self.benchmark_status_label.setText)
        self.benchmark_worker.benchmark_finished.connect(self._handle_benchmark_finished)
        self.benchmark_worker.error_occurred.connect(self._handle_benchmark_error)
        self.benchmark_worker.start()

    def _benchmark_source(self):
        if self.current_mode == DetectionMode.IMAGE:
            return self.current_image_path
        if self.current_mode == DetectionMode.FOLDER:
            return self.current_folder_images
        if self.current_mode == DetectionMode.VIDEO:
            return self.current_video_path
        return None

    def _handle_benchmark_finished(self, rows: List[Dict]) -> None:
        self.benchmark_rows = rows
        self.benchmark_worker = None
        self.benchmark_run_button.setEnabled(True)
        self.benchmark_export_button.setEnabled(bool(rows))
        self._populate_benchmark_table()
        if rows:
            best = max(rows, key=lambda item: item["fps"])
            self.benchmark_status_label.setText(f"参数对比完成：最快尺寸 {best['imgsz']}，FPS {best['fps']:.1f}")
        else:
            self.benchmark_status_label.setText("参数对比已取消。")
        self._refresh_current_analysis_tab()

    def _handle_benchmark_error(self, message: str) -> None:
        self.benchmark_worker = None
        self.benchmark_run_button.setEnabled(True)
        self.benchmark_export_button.setEnabled(False)
        self.benchmark_status_label.setText("参数对比失败")
        QMessageBox.critical(self, "参数对比失败", message)

    def _populate_benchmark_table(self) -> None:
        if getattr(self, "benchmark_table", None) is None:
            return
        self.benchmark_table.setRowCount(len(self.benchmark_rows))
        for row_index, row in enumerate(self.benchmark_rows):
            values = [
                row["imgsz"],
                row["samples"],
                f"{row['elapsed_s']:.3f}",
                f"{row['fps']:.1f}",
                f"{row['avg_ms']:.2f}",
                row["boxes"],
                f"{row['avg_conf']:.3f}",
            ]
            for col, value in enumerate(values):
                self.benchmark_table.setItem(row_index, col, QTableWidgetItem(str(value)))
        if getattr(self, "benchmark_export_button", None) is not None:
            self.benchmark_export_button.setEnabled(bool(self.benchmark_rows))

    def export_benchmark_csv(self) -> None:
        if not self.benchmark_rows:
            QMessageBox.information(self, "提示", "当前没有可导出的参数对比结果。")
            return
        save_path, _ = QFileDialog.getSaveFileName(self, "导出参数对比", "benchmark_compare.csv", "CSV 文件 (*.csv)")
        if not save_path:
            return
        with open(save_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(
                f,
                fieldnames=["imgsz", "samples", "elapsed_s", "fps", "avg_ms", "boxes", "avg_conf", "runtime"],
            )
            writer.writeheader()
            writer.writerows(self.benchmark_rows)
        self.statusBar().showMessage(f"参数对比结果已导出到: {save_path}")

    def export_analysis_report(self) -> None:
        save_path, _ = QFileDialog.getSaveFileName(self, "生成实验报告", "detection_experiment_report.md", "Markdown 文件 (*.md)")
        if not save_path:
            return
        if not save_path.lower().endswith(".md"):
            save_path += ".md"
        with open(save_path, "w", encoding="utf-8-sig") as f:
            f.write(self._build_markdown_report())
        self.statusBar().showMessage(f"实验报告已生成: {save_path}")

    def export_word_report(self) -> None:
        save_path, _ = QFileDialog.getSaveFileName(self, "生成Word实验报告", "detection_experiment_report.docx", "Word 文件 (*.docx)")
        if not save_path:
            return
        if not save_path.lower().endswith(".docx"):
            save_path += ".docx"
        try:
            from docx import Document
            from docx.shared import Inches
        except Exception:
            self._save_basic_docx_report(save_path)
            self.statusBar().showMessage(f"已生成简版Word实验报告: {save_path}")
            return

        doc = Document()
        doc.add_heading("水下目标检测实验报告", level=1)
        doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph(f"当前模型：{Path(self.detector.model_path).name if self.detector.model_path else '未加载'}")
        doc.add_paragraph(f"当前后端：{self.detector.runtime_label() if self.detector.is_loaded else '未加载'}")
        doc.add_paragraph(f"数据来源：{self._current_source_text()}")
        doc.add_heading("当前实验统计", level=2)
        for line in self._build_current_stats_text().splitlines():
            doc.add_paragraph(line)

        chart_dir = Path(self.temp_root) / "report_charts"
        chart_dir.mkdir(parents=True, exist_ok=True)
        class_chart = chart_dir / "class_distribution.png"
        timeline_chart = chart_dir / "frame_timeline.png"
        self._build_class_distribution_chart().save(str(class_chart))
        self._build_timeline_chart().save(str(timeline_chart))
        doc.add_heading("统计图表", level=2)
        doc.add_paragraph("类别数量分布")
        doc.add_picture(str(class_chart), width=Inches(6.2))
        doc.add_paragraph("FPS 与目标数量帧级趋势")
        doc.add_picture(str(timeline_chart), width=Inches(6.2))

        if self.benchmark_rows:
            doc.add_heading("输入尺寸对比", level=2)
            table = doc.add_table(rows=1, cols=7)
            table.style = "Table Grid"
            headers = ["输入尺寸", "样本数", "总耗时(s)", "FPS", "平均耗时(ms)", "目标数", "平均置信度"]
            for idx, header in enumerate(headers):
                table.rows[0].cells[idx].text = header
            for row in self.benchmark_rows:
                cells = table.add_row().cells
                values = [
                    row["imgsz"],
                    row["samples"],
                    f"{row['elapsed_s']:.3f}",
                    f"{row['fps']:.1f}",
                    f"{row['avg_ms']:.2f}",
                    row["boxes"],
                    f"{row['avg_conf']:.3f}",
                ]
                for idx, value in enumerate(values):
                    cells[idx].text = str(value)

        if self.alert_records:
            doc.add_heading("异常帧记录", level=2)
            table = doc.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            for idx, header in enumerate(["来源", "帧序号", "异常说明"]):
                table.rows[0].cells[idx].text = header
            for alert in self.alert_records:
                cells = table.add_row().cells
                cells[0].text = str(alert.get("source", ""))
                cells[1].text = str(alert.get("frame_index", ""))
                cells[2].text = "；".join(alert.get("alerts", []))

        sessions = self.backend.list_sessions(limit=5)
        if sessions:
            doc.add_heading("最近实验记录", level=2)
            table = doc.add_table(rows=1, cols=6)
            table.style = "Table Grid"
            for idx, header in enumerate(["ID", "模式", "模型", "状态", "采样帧", "目标数"]):
                table.rows[0].cells[idx].text = header
            for session in sessions:
                cells = table.add_row().cells
                values = [
                    session.get("id", ""),
                    session.get("mode", ""),
                    session.get("model_name", ""),
                    session.get("status", ""),
                    session.get("frame_count", 0),
                    session.get("detection_count", 0),
                ]
                for idx, value in enumerate(values):
                    cells[idx].text = str(value)
        doc.save(save_path)
        self.statusBar().showMessage(f"Word实验报告已生成: {save_path}")

    def _save_basic_docx_report(self, save_path: str) -> None:
        """Generate a small DOCX with only the standard library when python-docx is unavailable."""
        import zipfile
        from xml.sax.saxutils import escape as xml_escape

        def text_xml(text) -> str:
            return xml_escape(str(text))

        def paragraph(text, *, bold: bool = False, size: int = 22) -> str:
            run_props = f"<w:rPr>{'<w:b/>' if bold else ''}<w:sz w:val=\"{size}\"/></w:rPr>"
            return f"<w:p><w:r>{run_props}<w:t>{text_xml(text)}</w:t></w:r></w:p>"

        def table(headers: List[str], rows: List[List[object]]) -> str:
            def cell(value, *, bold: bool = False) -> str:
                return (
                    "<w:tc><w:tcPr><w:tcW w:w=\"2200\" w:type=\"dxa\"/></w:tcPr>"
                    f"{paragraph(value, bold=bold, size=18)}</w:tc>"
                )

            header_xml = "".join(cell(header, bold=True) for header in headers)
            row_xml = [f"<w:tr>{header_xml}</w:tr>"]
            for row in rows:
                row_xml.append("<w:tr>" + "".join(cell(value) for value in row) + "</w:tr>")
            return "<w:tbl><w:tblPr><w:tblW w:w=\"0\" w:type=\"auto\"/></w:tblPr>" + "".join(row_xml) + "</w:tbl>"

        body_parts = [
            paragraph("水下目标检测实验报告", bold=True, size=34),
            paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"),
            paragraph(f"当前模型：{Path(self.detector.model_path).name if self.detector.model_path else '未加载'}"),
            paragraph(f"当前后端：{self.detector.runtime_label() if self.detector.is_loaded else '未加载'}"),
            paragraph(f"数据来源：{self._current_source_text()}"),
            paragraph("当前实验统计", bold=True, size=28),
        ]
        for line in self._build_current_stats_text().splitlines():
            body_parts.append(paragraph(line, size=20))

        if self.benchmark_rows:
            body_parts.append(paragraph("输入尺寸对比", bold=True, size=28))
            body_parts.append(
                table(
                    ["输入尺寸", "样本数", "总耗时(s)", "FPS", "平均耗时(ms)", "目标数", "平均置信度"],
                    [
                        [
                            row["imgsz"],
                            row["samples"],
                            f"{row['elapsed_s']:.3f}",
                            f"{row['fps']:.1f}",
                            f"{row['avg_ms']:.2f}",
                            row["boxes"],
                            f"{row['avg_conf']:.3f}",
                        ]
                        for row in self.benchmark_rows
                    ],
                )
            )

        if self.alert_records:
            body_parts.append(paragraph("异常帧记录", bold=True, size=28))
            body_parts.append(
                table(
                    ["来源", "帧序号", "异常说明"],
                    [
                        [alert.get("source", ""), alert.get("frame_index", ""), "；".join(alert.get("alerts", []))]
                        for alert in self.alert_records
                    ],
                )
            )

        sessions = self.backend.list_sessions(limit=5)
        if sessions:
            body_parts.append(paragraph("最近实验记录", bold=True, size=28))
            body_parts.append(
                table(
                    ["ID", "模式", "模型", "状态", "采样帧", "目标数"],
                    [
                        [
                            session.get("id", ""),
                            session.get("mode", ""),
                            session.get("model_name", ""),
                            session.get("status", ""),
                            session.get("frame_count", 0),
                            session.get("detection_count", 0),
                        ]
                        for session in sessions
                    ],
                )
            )

        document_xml = (
            "<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>"
            "<w:document xmlns:w=\"http://schemas.openxmlformats.org/wordprocessingml/2006/main\">"
            "<w:body>"
            + "".join(body_parts)
            + "<w:sectPr><w:pgSz w:w=\"11906\" w:h=\"16838\"/>"
            "<w:pgMar w:top=\"1440\" w:right=\"1440\" w:bottom=\"1440\" w:left=\"1440\" "
            "w:header=\"708\" w:footer=\"708\" w:gutter=\"0\"/></w:sectPr>"
            "</w:body></w:document>"
        )
        content_types = (
            "<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>"
            "<Types xmlns=\"http://schemas.openxmlformats.org/package/2006/content-types\">"
            "<Default Extension=\"rels\" ContentType=\"application/vnd.openxmlformats-package.relationships+xml\"/>"
            "<Default Extension=\"xml\" ContentType=\"application/xml\"/>"
            "<Override PartName=\"/word/document.xml\" "
            "ContentType=\"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml\"/>"
            "</Types>"
        )
        rels = (
            "<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>"
            "<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\">"
            "<Relationship Id=\"rId1\" "
            "Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument\" "
            "Target=\"word/document.xml\"/>"
            "</Relationships>"
        )

        with zipfile.ZipFile(save_path, "w", compression=zipfile.ZIP_DEFLATED) as docx_file:
            docx_file.writestr("[Content_Types].xml", content_types)
            docx_file.writestr("_rels/.rels", rels)
            docx_file.writestr("word/document.xml", document_xml)

    def _build_markdown_report(self) -> str:
        current_stats = self._build_current_stats_text()
        lines = [
            "# 水下目标检测实验报告",
            "",
            f"- 生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            f"- 当前模型：{Path(self.detector.model_path).name if self.detector.model_path else '未加载'}",
            f"- 当前后端：{self.detector.runtime_label() if self.detector.is_loaded else '未加载'}",
            f"- 输入尺寸：{self._current_imgsz()}",
            f"- 数据来源：{self._current_source_text()}",
            "",
            "## 当前实验统计",
            "",
            "```text",
            current_stats,
            "```",
            "",
        ]

        if self.benchmark_rows:
            lines.extend(
                [
                    "## 输入尺寸对比",
                    "",
                    "| 输入尺寸 | 样本数 | 总耗时(s) | FPS | 平均耗时(ms) | 目标数 | 平均置信度 |",
                    "|---:|---:|---:|---:|---:|---:|---:|",
                ]
            )
            for row in self.benchmark_rows:
                lines.append(
                    f"| {row['imgsz']} | {row['samples']} | {row['elapsed_s']:.3f} | "
                    f"{row['fps']:.1f} | {row['avg_ms']:.2f} | {row['boxes']} | {row['avg_conf']:.3f} |"
                )
            lines.append("")

        if self.alert_records:
            lines.extend(
                [
                    "## 异常帧记录",
                    "",
                    "| 来源 | 帧序号 | 异常说明 |",
                    "|---|---:|---|",
                ]
            )
            for alert in self.alert_records:
                lines.append(
                    f"| {alert.get('source', '')} | {alert.get('frame_index', '')} | "
                    f"{'；'.join(alert.get('alerts', []))} |"
                )
            lines.append("")

        sessions = self.backend.list_sessions(limit=5)
        if sessions:
            lines.extend(
                [
                    "## 最近实验记录",
                    "",
                    "| ID | 模式 | 模型 | 后端 | 尺寸 | 状态 | 采样帧 | 目标数 | 平均FPS |",
                    "|---:|---|---|---|---:|---|---:|---:|---:|",
                ]
            )
            for session in sessions:
                lines.append(
                    f"| {session.get('id', '')} | {session.get('mode', '')} | {session.get('model_name', '')} | "
                    f"{session.get('runtime', '')} | {session.get('imgsz', '')} | {session.get('status', '')} | "
                    f"{session.get('frame_count', 0)} | {session.get('detection_count', 0)} | "
                    f"{self._format_float(session.get('avg_fps'), 1)} |"
                )
        return "\n".join(lines) + "\n"

    def _current_source_text(self) -> str:
        if self.current_mode == DetectionMode.IMAGE:
            return self.current_image_path or "未选择"
        if self.current_mode == DetectionMode.FOLDER:
            return self.folder_path_label.toolTip() or self.folder_path_label.text()
        if self.current_mode == DetectionMode.VIDEO:
            return self.current_video_path or "未选择"
        if self.current_mode == DetectionMode.CAMERA:
            return f"camera:{self.current_camera_id}"
        return "未选择"

    def _format_float(self, value, digits: int = 2) -> str:
        if value is None:
            return "--"
        try:
            return f"{float(value):.{digits}f}"
        except (TypeError, ValueError):
            return "--"

    # ----------------------- Save / export -----------------------
    def handle_save_action(self) -> None:
        if self.current_mode == DetectionMode.CAMERA and self.worker:
            self._toggle_camera_recording()
            return
        self.save_results()

    def save_results(self) -> None:
        if self.current_mode == DetectionMode.IMAGE:
            if not os.path.exists(self.temp_image_path):
                QMessageBox.warning(self, "提示", "当前没有可保存的图片结果。")
                return
            save_path, _ = QFileDialog.getSaveFileName(self, "保存检测图片", "", "图片文件 (*.jpg *.png *.bmp)")
            if save_path:
                safe_copy(self.temp_image_path, save_path)
                self.statusBar().showMessage(f"检测图片已保存到: {save_path}")
        elif self.current_mode == DetectionMode.FOLDER:
            if self._folder_result_count() > 0:
                save_dir = QFileDialog.getExistingDirectory(self, "选择保存目录")
                if save_dir:
                    target = os.path.join(save_dir, "detected_images")
                    if os.path.exists(target):
                        shutil.rmtree(target, ignore_errors=True)
                    shutil.copytree(self.temp_folder_dir, target)
                    self.statusBar().showMessage(f"文件夹检测结果已保存到: {target}")
                return

            if self.last_annotated_image is None:
                QMessageBox.warning(self, "提示", "当前没有可保存的检测结果。")
                return
            save_path, _ = QFileDialog.getSaveFileName(
                self,
                "保存当前预览结果",
                "detected_preview.jpg",
                "图片文件 (*.jpg *.png *.bmp)",
            )
            if save_path:
                if not Path(save_path).suffix:
                    save_path += ".jpg"
                cv2.imwrite(save_path, self.last_annotated_image)
                self.statusBar().showMessage(f"当前预览结果已保存到: {save_path}")
        elif self.current_mode == DetectionMode.VIDEO:
            if not os.path.exists(self.temp_video_path):
                QMessageBox.warning(self, "提示", "当前没有可保存的视频结果。")
                return
            save_path, _ = QFileDialog.getSaveFileName(self, "保存检测视频", "", "视频文件 (*.mp4)")
            if save_path:
                safe_copy(self.temp_video_path, save_path)
                self.statusBar().showMessage(f"检测视频已保存到: {save_path}")

    def export_log(self) -> None:
        if not self.session_records and self.backend_session_id is None:
            QMessageBox.information(self, "提示", "当前没有可导出的检测日志。")
            return
        fmt, ok = QInputDialog.getItem(
            self,
            "选择导出格式",
            "请选择检测结果导出格式：",
            ["CSV检测日志", "JSON检测日志", "YOLO TXT标注目录", "COCO JSON标注"],
            0,
            False,
        )
        if not ok:
            return
        try:
            if fmt == "CSV检测日志":
                save_path, _ = QFileDialog.getSaveFileName(self, "导出检测日志", "detection_log.csv", "CSV 文件 (*.csv)")
                if not save_path:
                    return
                if self.backend_session_id is not None:
                    count = self.backend.export_detections_csv(self.backend_session_id, save_path)
                    self.statusBar().showMessage(f"后端检测日志已导出到: {save_path}（{count} 行）")
                else:
                    self._export_records_csv(save_path)
            elif fmt == "JSON检测日志":
                save_path, _ = QFileDialog.getSaveFileName(self, "导出JSON检测日志", "detection_log.json", "JSON 文件 (*.json)")
                if not save_path:
                    return
                self._export_records_json(save_path)
            elif fmt == "YOLO TXT标注目录":
                save_dir = QFileDialog.getExistingDirectory(self, "选择YOLO TXT保存目录")
                if not save_dir:
                    return
                count = self._export_records_yolo(save_dir)
                self.statusBar().showMessage(f"YOLO TXT标注已导出到: {save_dir}（{count} 个文件）")
            else:
                save_path, _ = QFileDialog.getSaveFileName(self, "导出COCO JSON", "detections_coco.json", "JSON 文件 (*.json)")
                if not save_path:
                    return
                self._export_records_coco(save_path)
        except Exception as exc:
            QMessageBox.critical(self, "导出失败", str(exc))

    def _export_records_csv(self, save_path: str) -> None:
        fieldnames = [
            "mode",
            "source",
            "frame_index",
            "target_index",
            "class_id",
            "class_name",
            "confidence",
            "bbox",
            "frame_width",
            "frame_height",
            "alerts",
        ]
        with open(save_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(self.session_records)
        self.statusBar().showMessage(f"日志已导出到: {save_path}")

    def _export_records_json(self, save_path: str) -> None:
        if not save_path.lower().endswith(".json"):
            save_path += ".json"
        payload = {
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "model": Path(self.detector.model_path).name if self.detector.model_path else "",
            "runtime": self.detector.runtime_label() if self.detector.is_loaded else "",
            "imgsz": self._current_imgsz(),
            "mode": self.current_mode.name.lower(),
            "source": self._current_source_text(),
            "records": self.session_records,
            "alerts": self.alert_records,
        }
        with open(save_path, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
        self.statusBar().showMessage(f"JSON日志已导出到: {save_path}")

    def _record_group_key(self, record: Dict) -> str:
        source = str(record.get("source", "source")).replace("\\", "_").replace("/", "_")
        frame = record.get("frame_index", 0)
        return f"{Path(source).stem}_frame_{frame}"

    def _export_records_yolo(self, save_dir: str) -> int:
        os.makedirs(save_dir, exist_ok=True)
        groups: Dict[str, List[Dict]] = {}
        for record in self.session_records:
            groups.setdefault(self._record_group_key(record), []).append(record)
        class_map = self._class_id_map()
        for key, records in groups.items():
            out_path = os.path.join(save_dir, f"{key}.txt")
            with open(out_path, "w", encoding="utf-8") as f:
                for record in records:
                    bbox = record.get("bbox") or []
                    fw = float(record.get("frame_width", 0) or 0)
                    fh = float(record.get("frame_height", 0) or 0)
                    class_id = int(record.get("class_id", class_map.get(record.get("class_name", ""), 0)) or 0)
                    conf = float(record.get("confidence", 0.0) or 0.0)
                    if len(bbox) == 4 and fw > 0 and fh > 0:
                        x1, y1, x2, y2 = [float(v) for v in bbox]
                        cx = ((x1 + x2) / 2.0) / fw
                        cy = ((y1 + y2) / 2.0) / fh
                        bw = abs(x2 - x1) / fw
                        bh = abs(y2 - y1) / fh
                        f.write(f"{class_id} {cx:.6f} {cy:.6f} {bw:.6f} {bh:.6f} {conf:.6f}\n")
                    else:
                        f.write(f"{class_id} 0.000000 0.000000 0.000000 0.000000 {conf:.6f}\n")
        return len(groups)

    def _export_records_coco(self, save_path: str) -> None:
        if not save_path.lower().endswith(".json"):
            save_path += ".json"
        class_map = self._class_id_map()
        images = []
        annotations = []
        image_ids: Dict[str, int] = {}
        ann_id = 1
        for record in self.session_records:
            key = self._record_group_key(record)
            if key not in image_ids:
                image_ids[key] = len(image_ids) + 1
                images.append(
                    {
                        "id": image_ids[key],
                        "file_name": key,
                        "width": int(record.get("frame_width", 0) or 0),
                        "height": int(record.get("frame_height", 0) or 0),
                    }
                )
            bbox = record.get("bbox") or []
            if len(bbox) == 4:
                x1, y1, x2, y2 = [float(v) for v in bbox]
                w = max(0.0, x2 - x1)
                h = max(0.0, y2 - y1)
            else:
                x1 = y1 = w = h = 0.0
            class_name = record.get("class_name", "")
            annotations.append(
                {
                    "id": ann_id,
                    "image_id": image_ids[key],
                    "category_id": int(record.get("class_id", class_map.get(class_name, 0)) or 0),
                    "bbox": [x1, y1, w, h],
                    "area": w * h,
                    "score": float(record.get("confidence", 0.0) or 0.0),
                    "iscrowd": 0,
                }
            )
            ann_id += 1
        categories = [{"id": cid, "name": name} for name, cid in sorted(class_map.items(), key=lambda item: item[1])]
        with open(save_path, "w", encoding="utf-8") as f:
            json.dump({"images": images, "annotations": annotations, "categories": categories}, f, ensure_ascii=False, indent=2)
        self.statusBar().showMessage(f"COCO JSON已导出到: {save_path}")

    def _class_id_map(self) -> Dict[str, int]:
        mapping: Dict[str, int] = {}
        for record in self.session_records:
            name = str(record.get("class_name", ""))
            class_id = int(record.get("class_id", -1) or -1)
            if class_id < 0:
                class_id = len(mapping)
            mapping.setdefault(name, class_id)
        return mapping

    def _toggle_camera_recording(self) -> None:
        if not self.worker or self.current_mode != DetectionMode.CAMERA:
            return
        if self.is_paused:
            QMessageBox.warning(self, "提示", "请先恢复检测，再进行录制。")
            return
        if not self.is_recording:
            if os.path.exists(self.temp_camera_record_path):
                os.remove(self.temp_camera_record_path)
            self.worker.start_camera_recording(self.temp_camera_record_path)
            self.is_recording = True
            self.save_button.setText("停止录制")
            self.statusBar().showMessage("开始录制摄像头检测结果")
        else:
            self.worker.stop_camera_recording()
            self.is_recording = False
            self.save_button.setText("开始录制")

    def reset_view(self) -> None:
        self._stop_worker_if_needed()
        self._stop_benchmark_if_needed()
        self.current_mode = DetectionMode.NONE
        self.current_image_path = None
        self.current_folder_images = []
        self.current_video_path = None
        self.current_camera_id = None
        self.last_original_image = None
        self.last_clean_annotated_image = None
        self.last_annotated_image = None
        self.last_rows = []
        self._reset_session_records()

        self.image_path_label.setText("未选择图片")
        self.folder_path_label.setText("未选择文件夹")
        self.video_path_label.setText("未选择视频")
        self.image_path_label.setToolTip("")
        self.folder_path_label.setToolTip("")
        self.video_path_label.setToolTip("")

        self.progress_bar.setValue(0)
        self.progress_bar.setFormat("准备就绪")
        self.summary_label.setText("当前暂无检测结果")
        self.result_table.setRowCount(0)
        self.state_source_value.setText("无")
        self.state_status_value.setText("空闲")
        self.box_count_card.setText("检测框数量: 0")
        self._set_fps_value(0.0)
        self.mode_card.setText("当前模式: 未选择")
        self.label_result.clear()
        self.label_original.clear()
        self.label_result.setText("检测图像")

        self.start_button.setEnabled(False)
        self.pause_button.setEnabled(False)
        self.pause_button.setText("暂停")
        self.save_button.setEnabled(False)
        self.export_button.setEnabled(False)
        self.save_button.setText("保存结果")

        detector_loaded = self.detector.is_loaded
        self.image_detect_button.setEnabled(detector_loaded)
        self.folder_detect_button.setEnabled(detector_loaded)
        self.video_detect_button.setEnabled(detector_loaded)
        self.camera_detect_button.setEnabled(detector_loaded)
        self.remote_stream_button.setEnabled(detector_loaded)
        self.statusBar().showMessage("界面已重置")

    # ----------------------- Helpers -----------------------
    def _current_imgsz(self) -> int:
        return int(self.imgsz_spin.value()) if hasattr(self, "imgsz_spin") else self.config.inference_imgsz

    def _set_fps_value(self, fps: float) -> None:
        text = f"FPS: {fps:.1f}"
        self.fps_card.setText(text)
        self.summary_fps_label.setText(text)

    def _build_summary(self, rows: List[Dict], source_name: str, frame_index: int) -> str:
        if not rows:
            return f"来源: {source_name}\n帧/序号: {frame_index}\n未检测到目标"
        counter = Counter(row["class_name"] for row in rows)
        top_classes = ", ".join(f"{cls}×{count}" for cls, count in counter.most_common(5))
        max_conf = max(row["confidence"] for row in rows)
        return (
            f"来源: {source_name}\n"
            f"帧/序号: {frame_index}\n"
            f"目标数量: {len(rows)}\n"
            f"最高置信度: {max_conf:.2f}\n"
            f"类别统计: {top_classes}"
        )

    def _build_summary_html(self, summary: str, alerts: List[str]) -> str:
        body = "<br/>".join(escape(line) for line in summary.splitlines())
        if not alerts:
            return body
        alert_text = "；".join(alerts)
        return (
            f"{body}<br/>"
            f"<span style='color:#c00000; font-weight:700;'>"
            f"异常/报警提示：{escape(alert_text)}"
            f"</span>"
        )

    def _fill_table(self, rows: List[Dict]) -> None:
        self.result_table.setRowCount(len(rows))
        for r, row in enumerate(rows):
            self.result_table.setItem(r, 0, QTableWidgetItem(str(row["index"])))
            self.result_table.setItem(r, 1, QTableWidgetItem(row["class_name"]))
            self.result_table.setItem(r, 2, QTableWidgetItem(f"{row['confidence']:.3f}"))
        if rows:
            self.result_table.scrollToTop()

    def _update_progress(self, current: int, total: int) -> None:
        if total > 0:
            percent = int(current * 100 / total)
            self.progress_bar.setValue(percent)
            self.progress_bar.setFormat(f"{current} / {total}")
        else:
            self.progress_bar.setValue(0)
            self.progress_bar.setFormat("实时流处理中")

    def _append_session_records(self, payload: Dict) -> None:
        frame_h, frame_w = payload["original"].shape[:2] if payload.get("original") is not None else (0, 0)
        self.frame_samples.append(
            {
                "source": payload["source_name"],
                "frame_index": payload["frame_index"],
                "fps": float(payload.get("fps", 0.0) or 0.0),
                "box_count": int(payload.get("box_count", 0) or 0),
                "rows": [dict(row) for row in payload.get("rows", [])],
                "alerts": list(payload.get("alerts", []) or []),
            }
        )
        if self.record_limit_reached:
            return
        for row in payload["rows"]:
            if len(self.session_records) >= self.config.max_detection_records:
                self.record_limit_reached = True
                self.statusBar().showMessage("日志记录达到上限，后续记录将不再继续追加")
                break
            self.session_records.append(
                {
                    "mode": self.current_mode.name.lower(),
                    "source": payload["source_name"],
                    "frame_index": payload["frame_index"],
                    "target_index": row["index"],
                    "class_id": row.get("class_id", -1),
                    "class_name": row["class_name"],
                    "confidence": f"{row['confidence']:.6f}",
                    "bbox": row.get("bbox", []),
                    "center": row.get("center", []),
                    "frame_width": frame_w,
                    "frame_height": frame_h,
                    "alerts": "；".join(payload.get("alerts", []) or []),
                }
            )
        self.state_records_value.setText(str(len(self.session_records)))

    def _reset_session_records(self) -> None:
        self.session_records = []
        self.frame_samples = []
        self.alert_records = []
        self.benchmark_rows = []
        self.last_box_count_for_alert = None
        self.empty_frame_streak = 0
        self.record_limit_reached = False
        self.state_records_value.setText("0")

    def _update_path_label(self, label: QLabel, path: str) -> None:
        label.setText(short_path(path, keep_parts=3))
        label.setToolTip(path)

    def _clear_other_paths(self, keep: str) -> None:
        if keep != "image":
            self.image_path_label.setText("未选择图片")
            self.image_path_label.setToolTip("")
        if keep != "folder":
            self.folder_path_label.setText("未选择文件夹")
            self.folder_path_label.setToolTip("")
        if keep != "video":
            self.video_path_label.setText("未选择视频")
            self.video_path_label.setToolTip("")
        if keep == "camera":
            self.image_path_label.setText("未选择图片")
            self.folder_path_label.setText("未选择文件夹")
            self.video_path_label.setText("未选择视频")
            self.image_path_label.setToolTip("")
            self.folder_path_label.setToolTip("")
            self.video_path_label.setToolTip("")

    def _detect_local_cameras(self) -> List[int]:
        available = []
        for index in range(5):
            cap = cv2.VideoCapture(index)
            if cap.isOpened():
                available.append(index)
                cap.release()
        return available

    def _stop_worker_if_needed(self) -> None:
        if self.worker and self.worker.isRunning():
            self.worker.stop()
            self.worker.wait(1500)
        self.worker = None
        self.is_running = False
        self.is_paused = False
        self.is_recording = False

    def _stop_benchmark_if_needed(self) -> None:
        if self.benchmark_worker and self.benchmark_worker.isRunning():
            self.benchmark_worker.stop()
            self.benchmark_worker.wait(1500)
        self.benchmark_worker = None

    def closeEvent(self, event) -> None:  # noqa: N802
        self._save_settings()
        self._stop_worker_if_needed()
        self._stop_benchmark_if_needed()
        self.backend.finish_session(self.backend_session_id, status="closed", message="窗口关闭")
        self.backend.close()
        super().closeEvent(event)

